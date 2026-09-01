# src/extractors/main_parser.py

import re
import os
import pdfplumber
from pathlib import Path
from typing import ClassVar, List, Optional
from dataclasses import dataclass, field
from datetime import datetime
import sys
import math
import pandas as pd
import shutil
import importlib
import traceback

PROJECT_ROOT = Path(__file__).resolve().parents[2]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from src.utils.app_paths import (
    ensure_runtime_layout,
    get_input_dir,
    get_reference_dir,
    get_resource_roots,
)

# Зашитые дефолты на случай, если справочные CSV в reference/ отсутствуют
# (например до первого копирования бандла). Основной, редактируемый на
# месте источник — reference/coeff_dogovor.csv и
# reference/integral_party_keywords.csv (см. _load_contract_coefficients()
# и FinalUnifiedParser._load_party_keywords() ниже).
_DEFAULT_CONTRACT_COEFFICIENTS = {
    "2026008065": 1.33,
    "2025031516": 1.228,
}

_DEFAULT_INTEGRAL_PARTY_KEYWORDS = (
    "услуги",
    "всп.раб",  # "Всп.работы", "Всп. работы", "Всп.раб." — варианты написания встречаются все
    "вспомогательныеработы",
    "калибровк",
    "тех.дежурство",
    "технологическоедежурство",
    "работакомпл.партии",
    "работапартии",
    "переезд",
    "дефект",  # дефектоскопия, "дефект.и толщин."
    "толщинометри",
    "спецкаб",  # работы спецкабелем — угол не влияет (эмпирически, реальные акты)
)

_contract_coefficients_cache: dict | None = None


def _load_contract_coefficients() -> dict:
    """Коэффициенты по номеру договора — редактируемый справочник
    reference/coeff_dogovor.csv (колонки contract_number, coefficient).
    Правка коэффициента/добавление нового договора — правка этого CSV в
    рантайм-папке (AppData), без изменения кода. Если файла нет —
    используются зашитые дефолты."""
    global _contract_coefficients_cache
    if _contract_coefficients_cache is not None:
        return _contract_coefficients_cache

    csv_path = get_reference_dir() / "coeff_dogovor.csv"
    if csv_path.exists():
        try:
            df = pd.read_csv(csv_path, dtype=str, encoding="utf-8-sig")
            result = {}
            for _, row in df.iterrows():
                number = str(row.iloc[0]).strip()
                if not number:
                    continue
                try:
                    coeff = float(str(row.iloc[1]).strip().replace(",", "."))
                except Exception:
                    continue
                result[number] = coeff
            if result:
                _contract_coefficients_cache = result
                return _contract_coefficients_cache
        except Exception:
            pass

    _contract_coefficients_cache = dict(_DEFAULT_CONTRACT_COEFFICIENTS)
    return _contract_coefficients_cache

@dataclass
class WellData:
    """Класс для хранения данных по скважине"""
    filename: str
    field: str
    order: str
    depth: str
    angle: str
    temperature: str
    volume: str
    spo: str
    vm_task: str
    vm_price: str
    vm_count: str
    vm_total: str
    vm_table_count: str
    volume_sum_page1: str
    qty_sum_page3: str
    page2_start: str
    page2_end: str
    start_date: str
    end_date: str
    duration_hours: float = 0.0
    ml_confidence: float = 0.0
    ml_method: str = ""
    ml_suggestions: dict = field(default_factory=dict)
    ml_applied: dict = field(default_factory=dict)
    ml_spo_source: str = ""
    rate_check_status: str = ""
    rate_check_details: List[str] = field(default_factory=list)
    # Построчная сверка интегрального коэффициента (по факт. строкам расценок
    # из акта) — считается однажды в parse_all, пока PDF уже открыт
    integral_row_status: str = ""
    integral_row_details: List[str] = field(default_factory=list)
    task_number: str = ""
    contract_number: str = ""
    contract_coeff_value: float | None = None
    well_number: str = ""
    bush: str = ""
    zayavka_check_status: str = ""
    zayavka_check_details: List[str] = field(default_factory=list)
    # Номер сопоставленной заявки (задача) — для реестра "Проверка
    # акт-нарядов", колонка "Заявка".
    matched_zayavka_task: str = ""
    # Поля для реестра "Проверка акт-нарядов" (выгрузка по шаблону
    # заказчика): номер и дата самого акт-наряда (не заказа/договора),
    # буквенно-цифровой номер договора для отображения (в отличие от
    # contract_number — тот хранит только числовой формат, нужный
    # исключительно для проверки коэффициента по reference/coeff_dogovor.csv)
    # и итоговая сумма к оплате.
    act_number: str = ""
    act_date: str = ""
    contract_number_display: str = ""
    total_cost: str = ""
    performed_tasks: str = ""
    # Номер договора со страницы "АКТ-ЗАКАЗ" (стр.2) — сверяется с
    # contract_number_display (титульный лист, стр.1).
    page2_contract_number: str = ""
    # Номер задачи из заявки, встроенной в тот же PDF (обычно стр.4) —
    # для колонки "Заявка" в реестре, когда заявка не приходит отдельным
    # файлом (в реальных пакетах — почти всегда так).
    embedded_zayavka_task: str = ""
    # Свободный комментарий подрядчика со страницы "АКТ" (нижеподписавшиеся
    # ... составили настоящий акт о том, что ...) — недоход, остановка
    # прибора, осмотр перфоратора и т.д. Появляется не в каждом акте.
    # Для колонки "Комментарии" в реестре "Проверка акт-нарядов".
    contractor_comment: str = ""
    # Путь к исходному PDF — нужен проверке километража (km_parser заново
    # открывает файл для доступа к сканам/таблице переездов на стр.1).
    # Раньше км/температура считались только в пакетном режиме отдельно
    # от остальных проверок; для показа в одиночной проверке акта
    # (тот же список, что СПО/расценки/etc.) нужен путь до файла.
    source_path: str = ""
    # "Всего выполнено" по расценке 348 ("Спуск или подъем скв. прибора
    # через лубр.") из таблицы "Прочие виды работ" на акт-заказе (стр.2) —
    # для сверки с той же расценкой в таблице расценок акт-наряда (стр.1).
    # Обе величины в одних единицах (100 м).
    page2_spo: str = ""
    # Результаты доп.проверок, вычисленных в parse_all из строк расценок
    # (барометрия@53, тех.дежурство >4ч, пересечение термометрии 200/500)
    check_results: dict = field(default_factory=dict)

    # Shared DataFrame кэш: один на весь класс, не пересоздаётся при каждом parse_all
    _vm_price_df_cache: ClassVar["pd.DataFrame | None"] = None
    
    def __post_init__(self):
        """Вычисляем продолжительность после создания объекта"""
        self.duration_hours = self._calculate_duration()
    
    def _calculate_duration(self) -> float:
        """Вычисляет продолжительность работ в часах"""
        try:
            start_dt = datetime.strptime(self.start_date, '%d.%m.%Y %H:%M')
            end_dt = datetime.strptime(self.end_date, '%d.%m.%Y %H:%M')
            duration = end_dt - start_dt
            return duration.total_seconds() / 3600  # Часы
        except:
            return 0.0
    
    # Единый реестр проверок: (ключ, геттер результата, подпись по умолчанию,
    # молчать ли при статусе "neutral"). И print_with_units(), и
    # get_check_summary() идут по одному и тому же списку — раньше это были
    # два независимых места, которые несколько раз забывали синхронизировать
    # при добавлении новой проверки.
    #
    # Каждый геттер возвращает dict вида:
    #   {"status": "ok"|"bad"|"neutral"|"conditional_ok",
    #    "label": <необязательная замена подписи>,
    #    "value_text": <необязательный текст вместо иконки по умолчанию>,
    #    "details": [<уже полностью отформатированные строки, с отступом>]}
    _CHECK_REGISTRY_BEFORE_ML: ClassVar[tuple] = (
        ("spo", lambda self: self._check_volume_spo(), "Отчет по СПО", False),
        ("vm_cost", lambda self: self._check_vm_cost(), "Цена ВМ", False),
        ("rate", lambda self: self._check_rate(), "Проверка расценок", False),
        ("volume_qty", lambda self: self._check_volume_qty(),
         "Сравнение объемов (стр.1) и кол-ва (стр.3)", False),
        ("page2_dates", lambda self: self._check_page2_dates(), "Сравнение дат (стр.2)", False),
        ("contract_number_page2", lambda self: self._check_contract_number_page2(),
         "Сравнение номера договора (стр.2)", False),
    )
    _CHECK_REGISTRY_AFTER_ML: ClassVar[tuple] = (
        ("integral", lambda self: self._check_integral(), "Интегральный коэффициент", False),
        ("contract_coeff", lambda self: self._check_contract_coefficient(),
         "Коэффициент по договору", False),
        ("barometry_task53", lambda self: self.check_results.get("barometry_task53", {"status": "neutral"}),
         "Барометрия при задаче №53", True),
        ("tech_duty", lambda self: self._prefix_details(self.check_results.get("tech_duty", {"status": "neutral"})),
         "Тех.дежурство >4ч (нужен акт)", True),
        ("thermometry_overlap",
         lambda self: self._prefix_details(self.check_results.get("thermometry_overlap", {"status": "neutral"})),
         "Термометрия 200/500 (пересечение)", True),
        ("interval_length",
         lambda self: self._prefix_details(self.check_results.get("interval_length", {"status": "neutral"})),
         "Интервал ≤100м (стр.1)", True),
        ("spo_zakaz",
         lambda self: self._prefix_details(self.check_results.get("spo_zakaz", {"status": "neutral"})),
         "СПО: акт-наряд vs акт-заказ", True),
        ("zayavka", lambda self: self._check_zayavka(), "Сверка с заявкой", True),
        ("temperature", lambda self: self._check_temperature(), "Температура", False),
        ("km", lambda self: self._check_km(), "Километраж", False),
    )

    _STATUS_ICONS: ClassVar[dict] = {"ok": "✅", "bad": "❌"}

    def _print_check(self, getter, default_label: str, silent_on_neutral: bool) -> None:
        result = getter(self)
        status = result.get("status", "neutral")
        if status == "neutral" and silent_on_neutral:
            return
        label = result.get("label", default_label)
        if "value_text" in result:
            value_text = result["value_text"]
        elif status == "neutral":
            value_text = "данных нет"
        else:
            value_text = self._STATUS_ICONS.get(status, status)
        print(f"  {label}: {value_text}")
        for line in result.get("details", []):
            print(line)

    def print_with_units(self):
        """Вывод с единицами измерения"""
        print(f"\n{self.filename}:")
        print(f"  Месторождение: {self.field}")
        print(f"  Номер заказа: {self.order}")
        print(f"  Глубина забоя: {self.depth} м")
        print(f"  Угол наклона: {self.angle}°")
        print(f"  Температура воздуха: {self.temperature}°C")
        print(f"  СПО: {self.volume}")
        print(f"  СПО (таблица): {self.spo}")
        print(f"  Начало работ на скв: {self.start_date}")
        print(f"  Окончание работ на скв: {self.end_date}")
        print(f"  Продолжительность работ: {self.duration_hours:.2f} часов")
        if self.task_number:
            print(f"  Номер задачи: {self.task_number}")
        for _key, getter, label, silent in self._CHECK_REGISTRY_BEFORE_ML:
            self._print_check(getter, label, silent)
        self._print_ml_info()
        for _key, getter, label, silent in self._CHECK_REGISTRY_AFTER_ML:
            self._print_check(getter, label, silent)

    def _check_rate(self) -> dict:
        status = self.rate_check_status or "neutral"
        return {"status": status, "details": [f"    {line}" for line in self.rate_check_details]}

    def _to_float(self, value: str) -> float | None:
        try:
            val = float(str(value).replace(',', '.'))
            if math.isnan(val):
                return None
            return val
        except Exception:
            return None

    def _format_ru(self, value: float | None) -> str:
        if value is None:
            return ""
        s = f"{value:,.2f}"
        s = s.replace(",", "X").replace(".", ",").replace("X", " ")
        return s

    def _check_volume_spo(self) -> dict:
        spo_act = self._to_float(self.volume)
        spo_table = self._to_float(self.spo)
        if spo_act is None or spo_table is None:
            return {"status": "neutral"}
        if abs(spo_act) < 1000 <= abs(spo_table):
            spo_act = spo_act * 100
        elif abs(spo_act) >= 1000 > abs(spo_table):
            spo_table = spo_table * 100

        def _round_half_up(value: float) -> int:
            if value >= 0:
                return int(math.floor(value + 0.5))
            return int(math.ceil(value - 0.5))

        match = _round_half_up(spo_act) == _round_half_up(spo_table)
        return {
            "status": "ok" if match else "bad",
            "details": [
                f"    СПО в акте = {spo_act:.2f}",
                f"    СПО в таблице = {spo_table:.2f}",
            ],
        }

    def _check_vm_cost(self) -> dict:
        task = self.vm_task
        price_raw = self._to_float(self.vm_price)
        count = self._to_float(self.vm_count)
        total = self._to_float(self.vm_total)
        table_count = self._to_float(self.vm_table_count)
        unit_price = price_raw

        if count is not None and count > 0:
            if total is not None and total > 0:
                unit_from_total = total / count
                if unit_price is None or unit_price > unit_from_total * 1.5:
                    unit_price = unit_from_total
            elif unit_price is not None and unit_price > 10000:
                unit_price = unit_price / count

        if not task or unit_price is None:
            return {"status": "neutral", "label": "Стоимость ВМ", "value_text": "данных нет"}

        max_price, min_price = self._get_vm_prices(task, unit_price)
        if max_price is None and min_price is None:
            return {"status": "neutral", "label": "Стоимость ВМ", "value_text": "не найдено в таблице"}

        price_ok = min_price is not None and math.isclose(unit_price, min_price, rel_tol=0.0, abs_tol=0.1)

        details: List[str] = []
        if price_raw is not None:
            details.append(f"    Цена в акте = {self._format_ru(price_raw)}")
        if count is not None:
            details.append(f"    Кол-во = {int(count)}")
        if total is not None:
            details.append(f"    Итого = {self._format_ru(total)}")
        if table_count is not None:
            count_match = math.isclose(count, table_count, rel_tol=0.0, abs_tol=0.1) if count is not None else False
            count_icon = "✅" if count_match else "❌"
            details.append(f"    Сравнение кол-ва: {count_icon}")
            details.append(f"      Кол-во в таблице = {int(table_count)}")
        if max_price is not None:
            details.append(f"    Стоимость 1 отв. при max плотности = {self._format_ru(max_price)}")
        if min_price is not None:
            details.append(f"    Стоимость 1 отв. при min плотности = {self._format_ru(min_price)}")

        return {"status": "ok" if price_ok else "bad", "label": "Цена ВМ", "details": details}

    def _check_volume_qty(self) -> dict:
        vol_sum = self._to_float(self.volume_sum_page1)
        qty_sum = self._to_float(self.qty_sum_page3)
        if vol_sum is None or qty_sum is None:
            return {"status": "neutral"}
        match = math.isclose(vol_sum, qty_sum, rel_tol=0.0, abs_tol=0.1)
        return {
            "status": "ok" if match else "bad",
            "details": [
                f"    Сумма объемов (стр.1) = {vol_sum:.2f}",
                f"    Сумма кол-ва (стр.3) = {qty_sum:.2f}",
            ],
        }

    def _check_page2_dates(self) -> dict:
        start1 = self.start_date
        end1 = self.end_date
        start2 = self.page2_start
        end2 = self.page2_end
        if not start2 or not end2 or "не найдено" in (start2, end2):
            return {"status": "conditional_ok", "value_text": "условно ✅"}
        match = (start1 == start2) and (end1 == end2)
        return {
            "status": "ok" if match else "bad",
            "details": [
                f"    Начало работ: {start1} / {start2}",
                f"    Окончание работ: {end1} / {end2}",
            ],
        }

    def _check_contract_number_page2(self) -> dict:
        """Номер договора: титульный лист (стр.1) vs АКТ-ЗАКАЗ (стр.2) —
        тот же принцип, что и сравнение дат выше."""
        page1_value = self.contract_number_display
        page2_value = self.page2_contract_number
        if not page1_value or not page2_value:
            return {"status": "conditional_ok", "value_text": "условно ✅"}
        match = page1_value == page2_value
        return {
            "status": "ok" if match else "bad",
            "details": [f"    Номер договора: {page1_value} / {page2_value}"],
        }

    def _check_integral(self) -> dict:
        if not self.integral_row_status:
            return {"status": "neutral"}
        return {
            "status": self.integral_row_status,
            "details": [f"    {line}" for line in self.integral_row_details],
        }

    def _check_contract_coefficient(self) -> dict:
        expected = _load_contract_coefficients().get(self.contract_number)
        if not self.contract_number or expected is None:
            return {"status": "neutral"}
        if self.contract_coeff_value is None:
            return {"status": "neutral"}
        match = math.isclose(self.contract_coeff_value, expected, rel_tol=0.0, abs_tol=0.01)
        return {
            "status": "ok" if match else "bad",
            "details": [
                f"    Договор №{self.contract_number}: ожидается {expected:.3f}, "
                f"в акте {self.contract_coeff_value:.3f}"
            ],
        }

    def _check_zayavka(self) -> dict:
        if not self.zayavka_check_status:
            return {"status": "neutral"}
        return {
            "status": self.zayavka_check_status,
            "details": [f"    {line}" for line in self.zayavka_check_details],
        }

    def _check_temperature(self) -> dict:
        """Сверка температуры воздуха в акте со справочником — раньше
        считалась только в пакетном Excel-отчёте (table_parser.
        get_temperature_status), отдельно от остальных проверок; здесь —
        тот же расчёт, чтобы попасть в общий список и при проверке
        одного акта (кнопка "Пуск"), а не только пакетом."""
        from src.extractors import table_parser

        result = table_parser.get_temperature_status(self)
        status = result.get("status", "neutral")
        if status != "bad":
            return {"status": status}
        return {
            "status": "bad",
            "details": [
                f"    в акте {result['pdf_temp']:.1f}°C, по отчёту {result['avg_temp']:.1f}°C "
                f"(разница {result['diff']:.1f}°C, допуск 5°C)"
            ],
        }

    def _check_km(self) -> dict:
        """Сверка километража (переезды 1/3 гр. дорог, бездорожье) со
        справочником — раньше считалась только в пакетном Excel-отчёте
        (km_parser.compute_km_report, заново открывает PDF), здесь —
        тот же расчёт для единого списка проверок одного акта.

        Подробности (куст, акт/отчёт по каждой категории) выводим ВСЕГДА,
        а не только при расхождении — иначе статус "✅" без единой цифры
        не даёт понять, что вообще сверялось и с каким значением
        справочника совпало."""
        if not self.source_path:
            return {"status": "neutral"}
        from pathlib import Path
        from src.extractors import km_parser

        try:
            result = km_parser.compute_km_report(Path(self.source_path))
        except Exception:
            return {"status": "neutral"}
        status = result.get("status", "neutral")
        if status == "neutral":
            return {"status": "neutral"}

        lines = [f"    {line}" for line in km_parser.format_km_details(result)]
        if status == "conditional_ok":
            lines.append("    Переезд на другой объект — см. отметку на акт-заказе (стр.2)")
            return {"status": "conditional_ok", "value_text": "условно ✅", "details": lines}
        return {"status": status, "details": lines}

    @staticmethod
    def _prefix_details(result: dict) -> dict:
        """Строки в check_results (барометрия/тех.дежурство/термометрия)
        хранятся без отступа — добавляем его только для единого принтера,
        не трогая исходно сохранённый dict."""
        if not result.get("details"):
            return result
        return {**result, "details": [f"    {line}" for line in result["details"]]}

    def _check_interval_length(self, rows: List[dict]) -> dict:
        """Интервал перфорации (до-от) на стр.1 не должен превышать 100м —
        проверено на реальных актах: расценка "Перфорация на кабеле с
        привязкой" всегда даёт интервал ровно 100 м, независимо от объёма
        (единица измерения "опер.", не "100м"). Больший интервал в одной
        строке обычно означает ошибку заполнения (фактически нужны две
        операции перфорации по 100м, а не одна) — переплата за интервал,
        не подтверждённый отдельной операцией.

        Ограничено строго строками перфорации по названию: многие другие
        расценки (термометрия, влагометрия, профиль притока и т.п.) на
        реальных актах законно охватывают весь ствол скважины интервалом
        в сотни-тысячи метров — это не ошибка, это их обычный режим
        работы, и применять к ним тот же порог 100м неверно."""
        threshold = 100.0
        flagged: List[str] = []
        checked = 0
        for row in rows:
            name = str(row.get("name", "")).strip()
            if "перфораци" not in name.lower():
                continue
            interval_from = row.get("interval_from")
            interval_to = row.get("interval_to")
            if not isinstance(interval_from, (int, float)) or not isinstance(interval_to, (int, float)):
                continue
            length = interval_to - interval_from
            if length <= 0:
                continue
            checked += 1
            if length > threshold + 0.1:
                name_short = name if len(name) <= 70 else f"{name[:67]}..."
                flagged.append(
                    f"❌ {name_short} | интервал {interval_from:g}-{interval_to:g} = {length:.1f} м > {threshold:.0f} м"
                )
        if checked == 0:
            return {"status": "neutral"}
        if flagged:
            return {"status": "bad", "details": flagged}
        return {"status": "ok", "details": [f"В норме (≤{threshold:.0f} м), строк перфорации: {checked}"]}

    def _check_spo_zakaz(self, rows: List[dict]) -> dict:
        """СПО (расценка 348 "Спуск или подъем скв. прибора через лубр.")
        должно совпадать между таблицей расценок акт-наряда (стр.1) и
        таблицей "Прочие виды работ" акт-заказа (стр.2) — обе величины
        в одних единицах (100 м), без эвристики ×100/÷100, которая
        нужна старой проверке "spo" (сверка с текстовой строкой на
        странице "Справка по зарегистрированному материалу").

        Если расценки 348 нет вовсе на одной из двух страниц, а на
        другой есть ненулевое значение — это и есть основной случай,
        который старая проверка пропускала молча (self.volume оставался
        пустым, и _check_volume_spo просто не срабатывал): СПО-работа
        зафиксирована как выполненная на одной странице, но не выставлена
        к оплате (или наоборот) на другой."""
        naryad_value = None
        for row in rows:
            if row.get("rate_number") == "348":
                naryad_value = row.get("volume")
                break

        zakaz_raw = str(self.page2_spo or "").strip()
        if not zakaz_raw:
            # Таблицу "Прочие виды работ" на акт-заказе не нашли/не
            # распознали для этого акта — недостаточно данных, а не
            # подтверждённое расхождение.
            return {"status": "neutral"}

        zakaz_value = self._to_float(zakaz_raw) or 0.0
        naryad_value = naryad_value if naryad_value is not None else 0.0

        if naryad_value == 0.0 and zakaz_value == 0.0:
            return {"status": "neutral"}

        match = abs(naryad_value - zakaz_value) < 0.015
        return {
            "status": "ok" if match else "bad",
            "details": [
                f"    СПО в акт-наряде (расц. 348, стр.1) = {naryad_value:.2f}",
                f"    СПО в акт-заказе (стр.2) = {zakaz_value:.2f}",
            ],
        }

    def _check_barometry_task53(self, rows: List[dict]) -> dict:
        task = str(self.task_number or "").strip()
        if task != "53":
            return {"status": "neutral"}
        found = any("барометр" in str(row.get("name", "")).lower() for row in rows)
        return {"status": "ok" if found else "bad"}

    def _check_tech_duty_hours(self, rows: List[dict]) -> dict:
        threshold_hours = 4.0
        flagged: List[str] = []
        checked = 0
        for row in rows:
            name = str(row.get("name", "")).strip()
            norm = name.lower()
            if "тех.дежурство" not in norm and "технологическое дежурство" not in norm:
                continue
            volume = row.get("volume")
            if not isinstance(volume, (int, float)):
                continue
            checked += 1
            if volume > threshold_hours + 0.01:
                name_short = name if len(name) <= 70 else f"{name[:67]}..."
                flagged.append(
                    f"❌ {name_short} | {volume:.2f} ч > {threshold_hours:.0f} ч — нужен акт о простое"
                )
        if checked == 0:
            return {"status": "neutral"}
        if flagged:
            return {"status": "bad", "details": flagged}
        return {"status": "ok", "details": [f"В норме (≤{threshold_hours:.0f} ч), строк проверено: {checked}"]}

    def _tech_duty_hours_from_rows(self, rows: List[dict]) -> float | None:
        """Часы тех.дежурства из уже надёжно распознанных строк расценок
        (не OCR) — реальные акты сокращают название по-разному ("Тех.деж-
        во", "Тех.дежурство", "Технологическое дежурство"), поэтому ищем
        "тех" и "деж" рядом, а не точную фразу целиком."""
        total = 0.0
        found = False
        for row in rows:
            norm = str(row.get("name", "")).strip().lower()
            if not re.search(r'тех[а-я]*\.?\s*деж', norm):
                continue
            volume = row.get("volume")
            if isinstance(volume, (int, float)) and volume > 0:
                total += volume
                found = True
        return total if found else None

    def _finalize_contractor_comment(self, rows: List[dict]) -> None:
        """Итоговый комментарий в духе колонки "Комментарии" реестра
        заказчика: там почти всегда либо "+" (замечаний нет — большинство
        строк), либо короткий факт вида "Тех.деж. 4ч." или "Стоянка на
        глубине 2330 м.", а не абзац текста целиком. Часы тех.дежурства
        берём из строк расценок (надёжнее, чем распознавание акта о
        простое по OCR) и добавляем к уже найденному нарративному факту
        (стоянка/иное), если он есть."""
        hours = self._tech_duty_hours_from_rows(rows)
        parts = []
        if hours:
            hours_str = str(int(hours)) if hours == int(hours) else f"{hours:.2f}".rstrip("0").rstrip(".")
            parts.append(f"Тех.деж. {hours_str}ч.")
        if self.contractor_comment:
            parts.append(self.contractor_comment)
        self.contractor_comment = " ".join(parts) if parts else "+"

    def _check_thermometry_overlap(self, rows: List[dict]) -> dict:
        scale_re = re.compile(r'термометри.*?[MМ]\s*1\s*:\s*(200|500)', re.IGNORECASE)
        intervals_200: List[tuple] = []
        intervals_500: List[tuple] = []
        for row in rows:
            name = str(row.get("name", ""))
            if "термометр" not in name.lower():
                continue
            match = scale_re.search(name)
            if not match:
                continue
            i_from = row.get("interval_from")
            i_to = row.get("interval_to")
            if not isinstance(i_from, (int, float)) or not isinstance(i_to, (int, float)):
                continue
            lo, hi = min(i_from, i_to), max(i_from, i_to)
            bucket = intervals_200 if match.group(1) == "200" else intervals_500
            bucket.append((lo, hi, name))

        if not intervals_200 or not intervals_500:
            return {"status": "neutral"}

        overlaps: List[str] = []
        for lo1, hi1, name1 in intervals_200:
            for lo2, hi2, name2 in intervals_500:
                if lo1 < hi2 and lo2 < hi1:
                    overlaps.append(
                        f"❌ Пересечение: {name1} [{lo1:.1f}-{hi1:.1f}] и {name2} [{lo2:.1f}-{hi2:.1f}]"
                    )

        if overlaps:
            return {"status": "bad", "details": overlaps}
        return {
            "status": "ok",
            "details": [f"Пересечений не найдено (М1:200: {len(intervals_200)}, М1:500: {len(intervals_500)})"],
        }

    def get_check_summary(self) -> dict:
        """Структурированные статусы всех проверок (ok/bad/neutral/conditional_ok)
        без печати — используется для выгрузки пакетного Excel-отчёта. Идёт
        по тому же реестру _CHECK_REGISTRY_*, что и print_with_units(), —
        поэтому эти два места больше не могут разойтись."""
        summary = {}
        for key, getter, _label, _silent in self._CHECK_REGISTRY_BEFORE_ML + self._CHECK_REGISTRY_AFTER_ML:
            summary[key] = getter(self)
        return summary

    def _print_ml_info(self):
        if not self.ml_applied:
            return
        if self.ml_method == "ml-lines":
            return
        label_map = {
            "field": "Месторождение",
            "order": "Номер заказа",
            "temperature": "Температура воздуха",
            "angle": "Угол наклона",
            "volume": "СПО",
            "start_date": "Начало работ (стр.1)",
            "end_date": "Окончание работ (стр.1)",
            "page2_start": "Начало работ (стр.2)",
            "page2_end": "Окончание работ (стр.2)",
            "volume_sum_page1": "Сумма объемов (стр.1)",
            "qty_sum_page3": "Сумма объемов (стр.3)",
            "vm_price": "Цена ВМ",
            "vm_count": "Кол-во ВМ",
        }
        print("  🤖 ML-коррекция:")
        if self.ml_method:
            print(f"    Метод: {self.ml_method}, доверие: {self.ml_confidence:.2f}")
        if self.ml_spo_source:
            print(f"    СПО взято из поля: {self.ml_spo_source}")
        for attr, (old, new) in self.ml_applied.items():
            label = label_map.get(attr, attr)
            old_val = old if old else "—"
            print(f"    {label}: {old_val} → {new}")
    def _get_vm_prices(self, task: str, unit_price: float | None = None):
        excel_path = get_reference_dir() / "Стоимость ВМ задачи (не удалять).xlsx"
        if not excel_path.exists():
            return None, None
        if WellData._vm_price_df_cache is None:
            WellData._vm_price_df_cache = pd.read_excel(
                excel_path, sheet_name="Прейск_2019", header=None
            )
        df = WellData._vm_price_df_cache
        best_row = None
        best_diff = None
        for i in range(2, df.shape[0]):
            task_cell = str(df.iloc[i, 1]).strip()
            if task_cell != task:
                continue
            values = []
            for col in (7, 8, 9, 10):
                val = self._to_float(df.iloc[i, col])
                if val is not None:
                    values.append(val)
            if not values:
                continue
            if unit_price is None:
                best_row = i
                break
            diff = min(abs(unit_price - v) for v in values)
            if best_diff is None or diff < best_diff:
                best_diff = diff
                best_row = i

        if best_row is None:
            return None, None

        values = []
        for col in (7, 8, 9, 10):
            val = self._to_float(df.iloc[best_row, col])
            if val is not None:
                values.append(val)
        if not values:
            return None, None
        return max(values), min(values)

class FinalUnifiedParser:
    """Финальный объединенный парсер ВСЕХ данных"""
    _rate_reference_cache: dict | None = None
    _party_keywords_cache: tuple | None = None
    _prayskurant_codes_cache: set | None = None
    
    def parse_all(self, pdf_path: str) -> WellData:
        """Парсит ВСЕ значения включая даты и возвращает объект WellData"""
        ensure_runtime_layout(copy_reference=True)
        try:
            with pdfplumber.open(pdf_path) as pdf:
                # Титульный лист и Приложение №1 ищем по содержимому, а не по
                # фиксированному номеру страницы: реальные акты сдвигают эти
                # разделы (лишний лист согласования, заявка/акт готовности,
                # склеенные в тот же PDF, или вовсе обратный порядок листов —
                # см. 000088_*.pdf и 12459_*.pdf среди реальных актов).
                header_idx = self._find_page_index(pdf, ("Выполнение комплекса ГИРС по Договору",))
                if header_idx is None:
                    header_idx = 0
                prilozhenie_idx = self._find_page_index(pdf, ("Приложение №1", "Приложение N1", "Приложение N°1"))

                text = pdf.pages[header_idx].extract_text() if pdf.pages else ""
                all_text = "\n".join(page.extract_text() or "" for page in pdf.pages)

                filename = os.path.basename(pdf_path)

                # Вычисляем vm_count один раз (используется дважды)
                vm_count_val = self._parse_vm_count(all_text)
                # OCR страницы со сканом «АКТ-ЗАКАЗ» вызывается один раз
                scan_candidates = self._resolve_scan_candidates(pdf, header_idx, prilozhenie_idx)
                page2_start, page2_end, scan_idx = self._parse_page2_start_end(pdf, scan_candidates)
                vm_table_page_idx = scan_idx if scan_idx is not None else (
                    scan_candidates[0] if scan_candidates else header_idx + 1
                )

                well_data = WellData(
                    filename=filename,
                    source_path=pdf_path,
                    field=self._parse_field(text),
                    order=self._parse_order(text, pdf_path),
                    depth=self._parse_depth(text),
                    angle=self._parse_angle(text),
                    temperature=self._parse_temperature(text),
                    volume=self._parse_volume(text),
                    spo=self._parse_spo(all_text, pdf),
                    vm_task=self._parse_vm_task(text),
                    vm_price=self._parse_vm_price(all_text),
                    vm_count=vm_count_val,
                    vm_total=self._parse_vm_total(all_text),
                    vm_table_count=self._parse_vm_table_count(pdf, vm_count_val, vm_table_page_idx),
                    volume_sum_page1=self._parse_volume_sum_page1(text),
                    qty_sum_page3=self._parse_qty_sum_page3(pdf, prilozhenie_idx),
                    page2_start=page2_start,
                    page2_end=page2_end,
                    start_date=self._parse_date(text, 'начало'),
                    end_date=self._parse_date(text, 'окончание')
                )

                well_data.task_number = self._parse_task_number(text)
                well_data.contract_number = self._parse_contract_number(text)
                well_data.contract_coeff_value = self._parse_contract_coefficient_from_act(text)
                well_data.well_number, well_data.bush = self._parse_well_and_bush(text)
                well_data.contract_number_display = self._parse_contract_number_display(text)
                well_data.act_number, well_data.act_date = self._parse_act_number_and_date(text)
                well_data.total_cost = self._parse_total_cost(text)
                well_data.performed_tasks = self._parse_performed_tasks(text)
                ocr_task, well_data.contractor_comment = self._parse_zayavka_and_comment(pdf)
                if self._task_number_plausible(ocr_task, well_data.performed_tasks):
                    well_data.embedded_zayavka_task = ocr_task
                if scan_idx is not None and scan_idx < len(pdf.pages):
                    well_data.page2_contract_number = self._parse_contract_number_page2(pdf.pages[scan_idx])
                    well_data.page2_spo = self._parse_spo_zakaz(pdf.pages[scan_idx])

                rate_rows = self._extract_rate_rows_page1(pdf, header_idx)
                rate_status, rate_details = self._check_rate_prices_from_rows(rate_rows)
                well_data.rate_check_status = rate_status
                well_data.rate_check_details = rate_details

                # Построчная сверка интегрального коэффициента — переиспользует
                # уже извлечённые строки расценок (та же таблица, что и выше),
                # PDF повторно не открывается.
                integral_status, integral_details = self._check_integral_rows(
                    rate_rows, well_data.temperature, well_data.angle
                )
                well_data.integral_row_status = integral_status
                well_data.integral_row_details = integral_details

                well_data.check_results["barometry_task53"] = well_data._check_barometry_task53(rate_rows)
                well_data.check_results["tech_duty"] = well_data._check_tech_duty_hours(rate_rows)
                well_data.check_results["thermometry_overlap"] = well_data._check_thermometry_overlap(rate_rows)
                well_data.check_results["interval_length"] = well_data._check_interval_length(rate_rows)
                well_data.check_results["spo_zakaz"] = well_data._check_spo_zakaz(rate_rows)
                well_data._finalize_contractor_comment(rate_rows)

                if well_data.start_date == "не найдено" or well_data.end_date == "не найдено":
                    start_ocr, end_ocr = self._parse_page1_dates_ocr(pdf, header_idx)
                    if well_data.start_date == "не найдено" and start_ocr:
                        well_data.start_date = start_ocr
                    if well_data.end_date == "не найдено" and end_ocr:
                        well_data.end_date = end_ocr

                try:
                    from src.extractors.ml_assist import apply_ml_fallback
                except Exception:
                    try:
                        from ml_assist import apply_ml_fallback
                    except Exception:
                        apply_ml_fallback = None

                try:
                    if apply_ml_fallback is not None:
                        apply_ml_fallback(well_data, pdf_path=pdf_path)
                except Exception:
                    pass

                return well_data
                
        except Exception:
            if os.environ.get("AKT_DEBUG_PARSE_ERRORS") == "1":
                traceback.print_exc()
            return self._get_error_result(pdf_path)

    def _normalize_for_match(self, value: object) -> str:
        text = str(value) if value is not None else ""
        translit = {
            "A": "А", "B": "В", "C": "С", "E": "Е", "H": "Н", "K": "К",
            "M": "М", "O": "О", "P": "Р", "T": "Т", "X": "Х", "Y": "У",
            "a": "а", "b": "в", "c": "с", "e": "е", "h": "н", "k": "к",
            "m": "м", "o": "о", "p": "р", "t": "т", "x": "х", "y": "у",
        }
        text = "".join(translit.get(ch, ch) for ch in text)
        return re.sub(r"\s+", " ", text).strip().lower()

    def _parse_decimal(self, value: object) -> float | None:
        if value is None:
            return None
        text = str(value).replace("\u00a0", " ").strip()
        if not text:
            return None
        text = re.sub(r"[^0-9,.\- ]", "", text).replace(" ", "")
        if not text:
            return None
        if text.count(",") > 1 and "." not in text:
            parts = text.split(",")
            text = "".join(parts[:-1]) + "." + parts[-1]
        else:
            text = text.replace(",", ".")
        if text.count(".") > 1:
            parts = text.split(".")
            text = "".join(parts[:-1]) + "." + parts[-1]
        try:
            parsed = float(text)
            if math.isnan(parsed):
                return None
            return parsed
        except Exception:
            return None

    def _normalize_rate_number(self, value: object) -> str:
        if value is None:
            return ""
        text = str(value).strip()
        text = (
            text.replace("O", "0")
            .replace("О", "0")
            .replace("o", "0")
            .replace("о", "0")
        )
        digits = re.sub(r"\D", "", text)
        if not digits:
            return ""
        try:
            return str(int(digits))
        except Exception:
            return ""

    def _format_ru_decimal(self, value: float) -> str:
        return f"{value:,.2f}".replace(",", "X").replace(".", ",").replace("X", " ")

    def _format_rate_number(self, number: str) -> str:
        if not number:
            return ""
        return number.zfill(3) if len(number) <= 3 else number

    def _resolve_rate_csv_path(self) -> Path | None:
        reference_dir = get_reference_dir()
        preferred = reference_dir / "tabale_rascenki.csv"
        if preferred.exists():
            return preferred
        return None

    def _load_rate_reference(self) -> dict:
        if self._rate_reference_cache is not None:
            return self._rate_reference_cache
        csv_path = self._resolve_rate_csv_path()
        if not csv_path or not csv_path.exists():
            self._rate_reference_cache = {}
            return self._rate_reference_cache
        try:
            df = pd.read_csv(csv_path, dtype=str, encoding="utf-8-sig")
        except Exception:
            self._rate_reference_cache = {}
            return self._rate_reference_cache
        if df.empty:
            self._rate_reference_cache = {}
            return self._rate_reference_cache

        columns = list(df.columns)
        number_col = None
        price_col = None
        norm_col = None

        for col in columns:
            col_norm = self._normalize_for_match(str(col)).replace("_", " ")
            if number_col is None and "номер" in col_norm and "расцен" in col_norm:
                number_col = col
            if price_col is None and (
                ("стоим" in col_norm and ("руб" in col_norm or "расцен" in col_norm or "цена" in col_norm))
                or col_norm == "стоимость руб"
            ):
                price_col = col
            if norm_col is None and "норма" in col_norm and "врем" in col_norm:
                norm_col = col

        if number_col is None:
            number_col = columns[0]
        if price_col is None and len(columns) > 1:
            price_col = columns[1]

        rates: dict[str, dict[str, set[float]]] = {}
        for _, row in df.iterrows():
            number = self._normalize_rate_number(row.get(number_col))
            if not number:
                continue

            if number not in rates:
                rates[number] = {"prices": set(), "norms": set()}

            if price_col is not None:
                price = self._parse_decimal(row.get(price_col))
                if price is not None:
                    rates[number]["prices"].add(round(price, 2))

            if norm_col is not None:
                norm_value = self._parse_decimal(row.get(norm_col))
                if norm_value is not None:
                    rates[number]["norms"].add(round(norm_value, 2))

        self._rate_reference_cache = {
            key: {
                "prices": sorted(values.get("prices", set())),
                "norms": sorted(values.get("norms", set())),
            }
            for key, values in rates.items()
        }
        return self._rate_reference_cache

    def _load_prayskurant_codes(self) -> set:
        """Номера расценок из отдельного листа "Лист1" справочника по
        километражу ("17. Отчет по километражу.xlsx") — это "Прейскурант
        цен на отдельные виды геофизических исследований и работ в
        скважинах не учтенных в ЕНВиР-99-Л-ЗС" (мобилизация/демобилизация,
        переезды спецтехники, ГНКТ, X-MAC и т.д. — единичные услуги по
        договорным ценам). Сам файл прямо оговаривает (строки 198, 200
        листа): "коэффициенты, применяемые для ЕНВиР-99-Л-ЗС, не
        распространяют своё действие на расценки из настоящего
        прейскуранта" — то есть для этих расценок интегральный
        коэффициент в акте должен быть строго 1.00, а не посчитанным по
        температуре/углу, как для обычных строк ЕНВиР."""
        if self._prayskurant_codes_cache is not None:
            return self._prayskurant_codes_cache
        codes: set = set()
        try:
            import openpyxl

            excel_path = get_reference_dir() / "17. Отчет по километражу.xlsx"
            if excel_path.exists():
                wb = openpyxl.load_workbook(excel_path, data_only=True, read_only=True)
                if "Лист1" in wb.sheetnames:
                    ws = wb["Лист1"]
                    for row in ws.iter_rows(values_only=True):
                        value = row[2] if len(row) > 2 else None
                        if value is None:
                            continue
                        text = str(value).strip()
                        # Только настоящие номера расценок ("1807", "1833.1")
                        # — заголовки разделов и примечания текстовые, этому
                        # шаблону не соответствуют.
                        if re.fullmatch(r"\d+(\.\d+)?", text):
                            number = self._normalize_rate_number(text)
                            if number:
                                codes.add(number)
        except Exception:
            codes = set()
        self._prayskurant_codes_cache = codes
        return codes

    def _load_party_keywords(self) -> tuple:
        """Ключевые слова для классификации строк расценок (влияет ли угол
        наклона на интегральный коэффициент строки) — редактируемый
        справочник reference/integral_party_keywords.csv. Если файла нет
        (например на чистой установке до первого копирования бандла) —
        используется зашитый по умолчанию список."""
        if self._party_keywords_cache is not None:
            return self._party_keywords_cache

        csv_path = get_reference_dir() / "integral_party_keywords.csv"
        if csv_path.exists():
            try:
                df = pd.read_csv(csv_path, dtype=str, encoding="utf-8-sig")
                keywords = tuple(
                    str(value).strip().lower()
                    for value in df.iloc[:, 0].tolist()
                    if str(value).strip()
                )
                if keywords:
                    self._party_keywords_cache = keywords
                    return self._party_keywords_cache
            except Exception:
                pass

        self._party_keywords_cache = _DEFAULT_INTEGRAL_PARTY_KEYWORDS
        return self._party_keywords_cache

    def _extract_rate_rows_from_table(self, table: list) -> List[dict]:
        if not table:
            return []

        header_idx = None
        name_col = 0
        number_col = None
        price_col = None
        volume_col = None
        norm_col = None
        coeff_col = None
        spent_col = None
        interval_from_col = None
        interval_to_col = None

        for idx, row in enumerate(table[:8]):
            if not row:
                continue
            for col_idx, cell in enumerate(row):
                norm = self._normalize_for_match(cell)
                if "наименование" in norm and "работ" in norm:
                    name_col = col_idx
                if "номер" in norm and "расц" in norm:
                    number_col = col_idx
                if "расценка" in norm:
                    price_col = col_idx
                if "объем" in norm and "работ" in norm:
                    volume_col = col_idx
                if "норма" in norm and "времен" in norm:
                    norm_col = col_idx
                if "интег" in norm or "коэфиц" in norm:
                    coeff_col = col_idx
                if "затрат" in norm and "времен" in norm:
                    spent_col = col_idx
            if number_col is not None and price_col is not None:
                header_idx = idx
                # "интервал" ищем только в строке, уже подтверждённой как
                # заголовок таблицы (где нашлись номер расценки и цена) —
                # иначе случайное упоминание слова в титульной "шапке" акта
                # (например "интервалов перфорации" в описании задачи №53,
                # это одна огромная ячейка в table[0]) ложно фиксирует
                # interval_from_col на колонке названия работ, а
                # interval_to_col — на колонке номера расценки.
                for col_idx, cell in enumerate(row):
                    norm = self._normalize_for_match(cell)
                    if "интервал" in norm:
                        interval_from_col = col_idx
                        interval_to_col = col_idx + 1
                        break
                break

        if header_idx is None:
            header_blob = " ".join(
                self._normalize_for_match(cell)
                for row in table[:3]
                for cell in (row or [])
                if cell
            )
            if "наименование" in header_blob and "расценка" in header_blob:
                header_idx = 1 if len(table) > 1 else 0
                number_col = 1
                price_col = 7
                name_col = 0
                volume_col = 5
                norm_col = 6
                coeff_col = 8
                spent_col = 9
                interval_from_col = 3
                interval_to_col = 4
            else:
                return []

        rows: List[dict] = []
        stop_markers = (
            "итого по акт",
            "и т о г о по акт",
            "всего с учетом коэффициента",
            "стоимость вм",
            "всего к оплате",
        )

        for row in table[header_idx + 1:]:
            if not row:
                continue
            row_text = " ".join(str(cell).strip() for cell in row if cell)
            if not row_text:
                continue

            norm_row = self._normalize_for_match(row_text)
            if any(marker in norm_row for marker in stop_markers):
                break
            if "наименование работ" in norm_row:
                continue

            if number_col is None or number_col >= len(row):
                continue
            if price_col is None or price_col >= len(row):
                continue

            rate_number = self._normalize_rate_number(row[number_col])
            rate_price = self._parse_decimal(row[price_col])
            if not rate_number or rate_price is None:
                continue

            volume_value = None
            if volume_col is not None and volume_col < len(row):
                volume_value = self._parse_decimal(row[volume_col])

            norm_value = None
            if norm_col is not None and norm_col < len(row):
                norm_value = self._parse_decimal(row[norm_col])

            coeff_value = None
            if coeff_col is not None and coeff_col < len(row):
                coeff_value = self._parse_decimal(row[coeff_col])

            spent_value = None
            if spent_col is not None and spent_col < len(row):
                spent_value = self._parse_decimal(row[spent_col])

            interval_from_value = None
            if interval_from_col is not None and interval_from_col < len(row):
                interval_from_value = self._parse_decimal(row[interval_from_col])

            interval_to_value = None
            if interval_to_col is not None and interval_to_col < len(row):
                interval_to_value = self._parse_decimal(row[interval_to_col])

            name = ""
            if name_col < len(row) and row[name_col]:
                name = " ".join(str(row[name_col]).split())
            if not name:
                name = " ".join(str(cell).strip() for cell in row if cell)
                name = re.sub(r"\s+", " ", name)

            rows.append(
                {
                    "name": name,
                    "rate_number": rate_number,
                    "rate_price": round(rate_price, 2),
                    "volume": round(volume_value, 2) if volume_value is not None else None,
                    "norm_time": round(norm_value, 2) if norm_value is not None else None,
                    "integral_coeff": round(coeff_value, 4) if coeff_value is not None else None,
                    "spent_time": round(spent_value, 2) if spent_value is not None else None,
                    "interval_from": round(interval_from_value, 2) if interval_from_value is not None else None,
                    "interval_to": round(interval_to_value, 2) if interval_to_value is not None else None,
                }
            )

        return rows

    def _extract_rate_rows_page1(self, pdf, page_idx: int = 0) -> List[dict]:
        if not pdf.pages or page_idx >= len(pdf.pages):
            return []
        page = pdf.pages[page_idx]
        tables = page.extract_tables() or []
        rows: List[dict] = []
        for table in tables:
            rows.extend(self._extract_rate_rows_from_table(table))

        unique_rows: List[dict] = []
        seen: set[tuple] = set()
        for row in rows:
            key = (
                row.get("name", ""),
                row.get("rate_number", ""),
                row.get("rate_price", 0.0),
                row.get("norm_time"),
                row.get("spent_time"),
            )
            if key in seen:
                continue
            seen.add(key)
            unique_rows.append(row)
        return unique_rows

    def _check_rate_prices(self, pdf) -> tuple[str, List[str]]:
        rows = self._extract_rate_rows_page1(pdf)
        return self._check_rate_prices_from_rows(rows)

    def _check_rate_prices_from_rows(self, rows: List[dict]) -> tuple[str, List[str]]:
        reference = self._load_rate_reference()
        if not reference:
            return "neutral", []

        if not rows:
            return "neutral", []

        details: List[str] = []
        checked_rows = 0
        row_ok = 0
        row_bad = 0

        price_ok = 0
        price_bad = 0
        price_missing = 0

        norm_ok = 0
        norm_bad = 0
        norm_missing_ref = 0
        norm_missing_act = 0
        norm_enabled = any((entry.get("norms") or []) for entry in reference.values())

        spent_ok = 0
        spent_bad = 0
        spent_no_data = 0

        for row in rows:
            name = str(row.get("name", "")).strip()
            rate_number = str(row.get("rate_number", ""))
            rate_price = row.get("rate_price")
            if not rate_number:
                continue
            if not isinstance(rate_price, (int, float)):
                continue

            checked_rows += 1

            ref_entry = reference.get(rate_number, {})
            ref_prices = ref_entry.get("prices", []) if isinstance(ref_entry, dict) else []
            ref_norms = ref_entry.get("norms", []) if isinstance(ref_entry, dict) else []
            number_label = self._format_rate_number(rate_number)
            act_price_text = self._format_ru_decimal(float(rate_price))
            name_short = name if len(name) <= 80 else f"{name[:77]}..."

            row_has_issue = False

            if not ref_prices:
                price_missing += 1
                row_has_issue = True
                price_part = f"расценка: акт {act_price_text} / прил.: не найдено"
            else:
                price_match = any(
                    math.isclose(float(rate_price), float(ref), rel_tol=0.0, abs_tol=0.01)
                    for ref in ref_prices
                )
                ref_price_text = " / ".join(self._format_ru_decimal(float(v)) for v in ref_prices[:3])
                if len(ref_prices) > 3:
                    ref_price_text += " / ..."
                price_part = f"расценка: акт {act_price_text} / прил.: {ref_price_text}"
                if price_match:
                    price_ok += 1
                else:
                    price_bad += 1
                    row_has_issue = True

            act_norm = row.get("norm_time")
            norm_part = "норма: нет данных"
            if norm_enabled:
                if not ref_norms:
                    # Не штрафуем строку за отсутствие нормы в справочнике —
                    # у "договорных" расценок из Прейскуранта цен (в отличие
                    # от ЕНВиР) нормы времени в принципе не предусмотрено,
                    # и цена там уже сверяется отдельно (price_ok/price_bad).
                    norm_missing_ref += 1
                    norm_part = "норма: не предусмотрена для этой расценки"
                elif not isinstance(act_norm, (int, float)):
                    norm_missing_act += 1
                    row_has_issue = True
                    norm_part = "норма: в акте не найдена"
                else:
                    norm_match = any(
                        math.isclose(float(act_norm), float(ref), rel_tol=0.0, abs_tol=0.01)
                        for ref in ref_norms
                    )
                    ref_norm_text = " / ".join(self._format_ru_decimal(float(v)) for v in ref_norms[:3])
                    if len(ref_norms) > 3:
                        ref_norm_text += " / ..."
                    norm_part = f"норма: акт {self._format_ru_decimal(float(act_norm))} / прил.: {ref_norm_text}"
                    if norm_match:
                        norm_ok += 1
                    else:
                        norm_bad += 1
                        row_has_issue = True
            elif isinstance(act_norm, (int, float)):
                norm_part = f"норма: акт {self._format_ru_decimal(float(act_norm))}"
            else:
                norm_part = "норма: колонка в прил. не задана"

            act_spent = row.get("spent_time")
            act_volume = row.get("volume")
            act_coeff = row.get("integral_coeff")
            spent_part = "затраты: нет данных"
            if isinstance(act_spent, (int, float)) and isinstance(act_norm, (int, float)):
                if isinstance(act_volume, (int, float)):
                    coeff = float(act_coeff) if isinstance(act_coeff, (int, float)) else 1.0
                    expected_spent = float(act_volume) * float(act_norm) * coeff
                    expected_label = f"расч. {self._format_ru_decimal(expected_spent)}"
                else:
                    expected_spent = float(act_norm)
                    expected_label = f"норма {self._format_ru_decimal(float(act_norm))}"
                tolerance = max(0.5, abs(expected_spent) * 0.03)
                spent_match = math.isclose(float(act_spent), expected_spent, rel_tol=0.0, abs_tol=tolerance)
                spent_part = (
                    f"затраты: акт {self._format_ru_decimal(float(act_spent))} / {expected_label}"
                )
                if spent_match:
                    spent_ok += 1
                else:
                    spent_bad += 1
                    row_has_issue = True
            else:
                spent_no_data += 1

            row_icon = "✅" if not row_has_issue else "❌"
            if row_has_issue:
                row_bad += 1
            else:
                row_ok += 1
            details.append(
                f"{row_icon} {name_short} | №{number_label} | {price_part} | {norm_part} | {spent_part}"
            )

        if checked_rows == 0:
            return "neutral", []

        summary_lines = [
            f"Проверено работ: {checked_rows}",
            f"Итог по работам: ✅ {row_ok}, ❌ {row_bad}",
            f"Расценка: ✅ {price_ok}, ❌ {price_bad}, не найдено в прил.: {price_missing}",
        ]
        if norm_enabled:
            summary_lines.append(
                f"Норма времени (прил.): ✅ {norm_ok}, ❌ {norm_bad}, не найдено в прил.: {norm_missing_ref}, нет в акте: {norm_missing_act}"
            )
        else:
            summary_lines.append("Норма времени (прил.): колонка не найдена в CSV")
        summary_lines.append(
            f"Затраты времени (акт/расчет): ✅ {spent_ok}, ❌ {spent_bad}, нет данных: {spent_no_data}"
        )

        status_is_bad = (
            row_bad > 0
            or price_bad > 0
            or price_missing > 0
            or spent_bad > 0
            or (norm_enabled and (norm_bad > 0 or norm_missing_act > 0))
        )
        status = "bad" if status_is_bad else "ok"
        return status, summary_lines + details

    def _classify_integral_work_type(self, name: str) -> str:
        norm = re.sub(r"\s+", "", name.lower())
        for keyword in self._load_party_keywords():
            if keyword in norm:
                return "party"
        return "well"

    def _check_integral_rows(self, rows: List[dict], temperature: str, angle: str) -> tuple[str, List[str]]:
        try:
            from integral import CoefficientTable
        except Exception:
            return "neutral", []

        try:
            temp = float(str(temperature).replace(",", "."))
        except Exception:
            return "neutral", []

        try:
            angle_val = float(str(angle).replace(",", "."))
        except Exception:
            angle_val = None

        table = CoefficientTable()
        k_party = table.get_temperature_coefficient(temp)
        k_well = table.get_integral_coefficient(temp, angle_val) if angle_val is not None else None
        prayskurant_codes = self._load_prayskurant_codes()

        details: List[str] = []
        ok_count = 0
        bad_count = 0

        for row in rows:
            actual = row.get("integral_coeff")
            if not isinstance(actual, (int, float)):
                continue

            name = str(row.get("name", "")).strip()
            rate_number = str(row.get("rate_number", ""))
            # Расценки из Прейскуранта (мобилизация/переезды спецтехники,
            # ГНКТ, X-MAC и т.п.) — договорные, коэффициент ЕНВиР к ним не
            # применяется, ожидаем строго 1.00 независимо от температуры/угла.
            if rate_number and rate_number in prayskurant_codes:
                expected = 1.0
                kind = "ед.усл."
            else:
                work_type = self._classify_integral_work_type(name)
                expected = k_well if work_type == "well" else k_party
                kind = "скв." if work_type == "well" else "парт."
            if expected is None:
                continue

            match = math.isclose(float(actual), float(expected), rel_tol=0.0, abs_tol=0.02)
            name_short = name if len(name) <= 70 else f"{name[:67]}..."
            icon = "✅" if match else "❌"
            details.append(
                f"{icon} {name_short} | акт={float(actual):.2f} / расчёт={float(expected):.2f} ({kind})"
            )
            if match:
                ok_count += 1
            else:
                bad_count += 1

        if ok_count == 0 and bad_count == 0:
            return "neutral", []

        status = "bad" if bad_count > 0 else "ok"
        summary = f"Совпало: {ok_count}, расхождений: {bad_count}"
        return status, [summary] + details

    # ===== ОСНОВНЫЕ ДАННЫЕ =====
    
    def _parse_field(self, text: str) -> str:
        """Месторождение — общая с km_parser.parse_field логика (та же
        сверка со справочником по километражу, где эти 10 названий полей
        и берутся), иначе простые и составные (через дефис) названия
        полей заново разъезжались бы по двум местам, как уже бывало."""
        from src.extractors import km_parser

        field = km_parser.parse_field(text)
        return field if field else "не найдено"
    
    def _parse_order(self, text: str, pdf_path: str) -> str:
        """Номер заказа"""
        match = re.search(r'Заказ\s*№[^\d]*([\d\.\s]{6,10})', text)
        if match:
            digits = re.sub(r'[^\d]', '', match.group(1))
            if 6 <= len(digits) <= 7:
                return digits
        
        filename = os.path.basename(pdf_path)
        match = re.search(r'_(\d{6,7})\.pdf$', filename)
        return match.group(1) if match else "не найдено"

    def _parse_task_number(self, text: str) -> str:
        """Номер задачи (например 53, 87(Р), 500.4, 58.158)"""
        match = re.search(r'Задача\s*№\s*([0-9]+(?:\.[0-9]+)?(?:\([А-Яа-я]+\))?)', text)
        return match.group(1) if match else ""

    _ZAYAVKA_TASK_PATTERNS: ClassVar[tuple] = (
        # "ПГИ № 80 (...)" — старый цифровой формат заявки.
        re.compile(r'ПГИ\s*№\s*([0-9]+(?:\.[0-9]+)?(?:\([А-Яа-я]+\))?)'),
        # "Цель (задача) и интервал исследований: № 24 Опр...."
        re.compile(r'Цель\s*\(задача\)[^\n]{0,60}?№\s*([0-9]+(?:\.[0-9]+)?(?:\([А-Яа-я]+\))?)'),
        # "Цель, вид, объем заказываемых работ: 87 Опр...." — без "№".
        re.compile(r'Цель,?\s*вид,?\s*объем[^\n]{0,40}?:\s*([0-9]+(?:\.[0-9]+)?(?:\([А-Яа-я]+\))?)'),
        # общий запасной вариант — как в заголовке акт-наряда/акт-заказа.
        re.compile(r'Задача\s*№\s*([0-9]+(?:\.[0-9]+)?(?:\([А-Яа-я]+\))?)'),
    )

    def _task_number_plausible(self, ocr_task: str, performed_tasks: str) -> bool:
        """OCR номера задачи со скана заявки достаточно ненадёжен (буквы
        путаются с цифрами, суффиксы вроде "(S)"/".134" теряются или
        дают лишнюю цифру: "35(S)" стабильно читается как "358"). Прежде
        чем доверять результату, сверяем его ведущие цифры с уже надёжно
        распознанной задачей из шапки акта (task_number/performed_tasks,
        чистый текст, без OCR) — если они не совпадают, значение отбрасываем
        (в отчёте вместо него встанет запасной вариант performed_tasks),
        а не показываем вероятно garbled цифры."""
        if not ocr_task:
            return False
        ocr_lead = re.match(r'\d+', ocr_task)
        if not ocr_lead:
            return False
        ocr_lead = ocr_lead.group()
        for task in performed_tasks.split("+"):
            performed_lead = re.match(r'\d+', task.strip())
            if performed_lead and performed_lead.group() == ocr_lead:
                return True
        return not performed_tasks

    def _match_zayavka_task(self, text: str) -> str:
        for pattern in self._ZAYAVKA_TASK_PATTERNS:
            match = pattern.search(text)
            if match:
                return match.group(1)
        return ""

    # Короткие, самодостаточные факты внутри длинного текста страницы
    # "АКТ" — извлекаем ИХ, а не весь абзац целиком (реестр заказчика на
    # 2696 реальных строках показывает: 37% комментариев — это просто "+"
    # (замечаний нет), ~41% — короткое "Тех.деж. N ч.", ~5% — короткое
    # "Стоянка на гл. N м", и лишь остаток — действительно свободный текст).
    _COMMENT_STOYANKA_RE: ClassVar = re.compile(
        r'сто[яй]нк[а-я]*(?:\s+(?:прибора|шаблона|магнита)[а-я]*)?'
        r'(?:[^.]{0,40}?(?:на\s+)?гл(?:убине)?\.?)?[^.]{0,25}?'
        r'[\d][\d\s,.]*\s*м(?:етр[а-я]*)?\b[^.]{0,60}',
        re.IGNORECASE,
    )
    _COMMENT_TECH_DUTY_RE: ClassVar = re.compile(
        # "составило/составила" НЕ обязательно — реальные акты пишут и
        # без него ("Тех.дежурство при компрессировании, 3,5 часа."),
        # раньше жёсткое требование этого глагола роняло совпадение и
        # уходило в сырой fallback ниже. "tex" (латиницей) — частая OCR-
        # ошибка распознавания "Тех" на сканах (напр. акт 13074: "® Tex.
        # дежурство при компрессировании, 3,5 часа.").
        r'(?:тех[а-я]*|tex)\.?\s*деж[а-я]*[^.]{0,60}?(\d+[.,]?\d*)[^.]{0,20}?час',
        re.IGNORECASE,
    )

    def _extract_contractor_comment(self, text: str) -> str:
        """Текст свободного комментария подрядчика со страницы вида
        "ПАО «КОГАЛЫМНЕФТЕГЕОФИЗИКА» / АКТ / Мы, нижеподписавшиеся: ...
        составили настоящий акт о том, что <дата> на скважине № X ...".
        Такая страница появляется не в каждом акте (только когда есть что
        отметить — недоход, остановка прибора, осмотр перфоратора после
        извлечения и т.д.), в отличие от рутинного "АКТ проверки готовности
        скважины" (у него тоже есть "нижеподписавшиеся", но сразу за
        "составили настоящий акт" идёт неизменное "о том, что нами
        проверена готовность" — исключаем именно эту формулировку).

        Из найденной страницы вырезаем КОРОТКИЙ ключевой факт (стоянка на
        такой-то глубине, тех.дежурство N часов), а не абзац целиком —
        так короче и ближе к тому, как реально заполняется колонка
        "Комментарии" в реестре заказчика.

        Если короткий факт не нашёлся — это либо рутинная страница без
        замечаний ("опрессовка устьевого оборудования", "осмотр
        перфоратора" без повреждений — проверено: 0 из 2696 реальных
        комментариев содержат "опресс"), либо страница какого-то ещё не
        встреченного нам формата (другой бланк осмотра перфоратора,
        простой и т.п.). В обоих случаях НЕ отдаём сырой обрубок текста —
        он часто ещё и с OCR-искажениями ("Шо £02", "< 02") и выглядит
        как мусор, а не как факт. Пусть выше по стеку сработает дефолт
        "+". Единственное исключение — явно обнаруженная поломка/отказ:
        тут "+" был бы неправдой, поэтому отдаём короткую пометку с
        началом абзаца, чтобы человек нашёл и прочитал акт вручную."""
        match = re.search(r'составили\s+настоящий\s+[Аа]кт', text, re.IGNORECASE)
        if not match:
            return ""
        tail = text[match.end():match.end() + 250]
        if re.search(r'проверена\s+готовност', tail, re.IGNORECASE):
            return ""
        body = text[match.start():].strip()
        body = re.sub(r'\s+', ' ', body)

        tech_duty_match = self._COMMENT_TECH_DUTY_RE.search(body)
        stoyanka_match = self._COMMENT_STOYANKA_RE.search(body)
        facts = []
        if stoyanka_match:
            facts.append(re.sub(r'\s+', ' ', stoyanka_match.group()).strip(" ,;"))
        if tech_duty_match:
            hours = tech_duty_match.group(1).replace(",", ".")
            facts.append(f"Тех.деж. {hours}ч.")
        if facts:
            return " ".join(facts)

        # Именно ПОЛОЖИТЕЛЬНАЯ формулировка обнаруженной проблемы, а не
        # голый корень "поврежд" — иначе штатное "не имеют явных
        # повреждений" (то есть повреждений НЕТ) ложно считается аномалией.
        has_anomaly = re.search(
            r'обнаружил[аои]?\s+(?:следующ\w*\s+)?(?:повреждени|дефект)|'
            r'имеются?\s+повреждени|отказ(?:ал)?|брак(?:ован)?|не\s+сработал|негерметичн',
            body, re.IGNORECASE,
        )
        if has_anomaly:
            return "⚠ см. акт: " + body[:150]
        return ""

    def _ocr_full_page_text(self, page, pytesseract_module) -> str:
        try:
            from PIL import ImageOps
            image = page.to_image(resolution=200).original
            gray = ImageOps.grayscale(image)
            gray = ImageOps.autocontrast(gray)
            return pytesseract_module.image_to_string(gray, lang="rus+eng", config="--psm 6", timeout=20)
        except Exception:
            return ""

    def _parse_zayavka_and_comment(self, pdf) -> tuple[str, str]:
        """Один проход по страницам акта — номер задачи из встроенной
        ЗАЯВКИ (лист "ЗАЯВКА на проведение промыслово-геофизических
        исследований скважин", обычно 4-й) и текст комментария подрядчика
        (страница "АКТ" с "нижеподписавшиеся"), оба поля — с OCR-фоллбэком:
        на реальных актах обе страницы чаще скан, чем цифровой текст (по
        выборке 86 актов OCR потребовался для заявки в ~65 из 82 найденных
        случаев). Общий проход экономит OCR — не открываем те же страницы
        дважды под каждое поле по отдельности."""
        task_number = ""
        comment_text = ""
        pytesseract_module = None
        # Комментарий подрядчика — почти всегда на последних листах пакета
        # (после Справки по зарегистрированному материалу, ближе к концу),
        # а не на фиксированном номере страницы: сами акты разной длины
        # (от 7 до 16+ листов среди реальных актов). Ищем в последних 8
        # страницах — так же дёшево на коротком акте, но не пропускает
        # комментарий в длинном, и не сканирует OCR'ом весь документ,
        # если комментария в акте нет вовсе.
        total_pages = len(pdf.pages)
        comment_min_idx = max(3, total_pages - 8)

        for idx, page in enumerate(pdf.pages):
            if task_number and comment_text:
                break
            try:
                text = page.extract_text() or ""
            except Exception:
                text = ""
            needs_ocr = len(text.strip()) < 30
            ocr_text = None

            if not task_number:
                if "заявка" in text.lower() or "на проведение" in text.lower():
                    task_number = self._match_zayavka_task(text)
                if not task_number and needs_ocr and idx <= 6:
                    if pytesseract_module is None:
                        try:
                            import pytesseract as pytesseract_module
                            if not self._configure_tesseract(pytesseract_module):
                                pytesseract_module = None
                        except Exception:
                            pytesseract_module = None
                    if pytesseract_module is not None:
                        ocr_text = self._ocr_full_page_text(page, pytesseract_module)
                        ocr_lower = ocr_text.lower()
                        if "заявка" in ocr_lower or "на проведение" in ocr_lower:
                            task_number = self._match_zayavka_task(ocr_text)

            if not comment_text:
                if "нижеподписавш" in text.lower():
                    comment_text = self._extract_contractor_comment(text)
                elif needs_ocr and comment_min_idx <= idx:
                    if ocr_text is None:
                        if pytesseract_module is None:
                            try:
                                import pytesseract as pytesseract_module
                                if not self._configure_tesseract(pytesseract_module):
                                    pytesseract_module = None
                            except Exception:
                                pytesseract_module = None
                        if pytesseract_module is not None:
                            ocr_text = self._ocr_full_page_text(page, pytesseract_module)
                    if ocr_text and "нижеподписавш" in ocr_text.lower():
                        comment_text = self._extract_contractor_comment(ocr_text)

        return task_number, comment_text

    def _parse_performed_tasks(self, text: str) -> str:
        """Все задачи, перечисленные в шапке акта ('Задача №54.1', 'Задача
        №53' — акт может выполнять сразу несколько), объединённые через
        '+' в порядке появления — колонка "Проведенный ГИС" в реестре
        "Проверка акт-нарядов" (в отличие от task_number, берущего только
        первую — она нужна для сверки с заявкой и спецкейса недохода)."""
        matches = re.findall(r'Задача\s*№\s*([0-9]+(?:\.[0-9]+)?(?:\([А-Яа-я]+\))?)', text)
        return "+".join(matches)

    def _parse_contract_number(self, text: str) -> str:
        """Номер договора (только числовой формат — учитываются лишь
        договоры из reference/coeff_dogovor.csv; старые буквенно-цифровые
        номера вроде '22С3286' для этой проверки не применимы)."""
        match = re.search(r'[Дд]оговор[а-я]*\s*№\s*([0-9]{6,})', text)
        return match.group(1) if match else ""

    def _parse_contract_number_display(self, text: str) -> str:
        """Номер договора в исходном виде (буквенно-цифровой, напр.
        '23С1816', '22С3286') — для реестра "Проверка акт-нарядов", в
        отличие от _parse_contract_number(), который берёт только числовой
        формат ради сверки коэффициента."""
        match = re.search(r'по Договору\s*№\s*([0-9A-ZА-Я]+)', text)
        return match.group(1) if match else ""

    def _parse_act_number_and_date(self, text: str) -> tuple[str, str]:
        """Номер и дата акт-наряда из строки 'АКТ - НАРЯД № 00079 от
        06.01.2026' (буква после дефиса иногда распознаётся как латинская
        H вместо кириллической Н). Ведущие нули в номере отбрасываются —
        так же, как они показаны в реестре заказчика ('00079' -> '79')."""
        match = re.search(r'АКТ\s*-\s*[HН]?АРЯД\s*№\s*(\d+)\s*от\s*(\d{2}\.\d{2}\.\d{4})', text)
        if not match:
            return "", ""
        number = str(int(match.group(1)))
        return number, match.group(2)

    def _parse_total_cost(self, text: str) -> str:
        """Итоговая сумма к оплате по акту ('Всего к оплате 170 372,13')."""
        match = re.search(r'Всего к оплате\s+([\d\s]+,\d+)', text)
        if not match:
            return ""
        value = match.group(1).replace(" ", "").replace(",", ".")
        return value

    def _parse_contract_coefficient_from_act(self, text: str) -> float | None:
        """Значение из строки 'Всего с учетом коэффициента' на 1 странице акта."""
        match = re.search(r'Всего с учетом коэффициента[^\d]*([\d]+[,.]\d+)', text)
        if not match:
            return None
        try:
            return float(match.group(1).replace(",", "."))
        except Exception:
            return None

    def _parse_well_and_bush(self, text: str) -> tuple[str, str]:
        """Номер скважины и куста из строки 'Номер скважины / куст.......1867 / 151'.
        И номер скважины, и номер куста может иметь буквенный суффикс
        ('1996Л', '35А', '90Б') — это часть самого номера, а не опечатка,
        и без учёта суффикса матч проваливался целиком (не находилось ни
        скважины, ни куста) либо куст обрезался до одних цифр.
        Заполнитель между "куст" и номером — не только точки/пробелы, но и
        символ многоточия "…" (используется в реальных актах), поэтому
        [^\\d]* (любые не-цифры) надёжнее, чем [.\\s]*."""
        match = re.search(
            r'Номер скважины\s*/\s*куст[^\d]*(\d+[А-Яа-яA-Za-z]*)\s*/\s*(\d+[А-Яа-яA-Za-z]*)', text
        )
        if match:
            return match.group(1), match.group(2)
        return "", ""

    def _parse_depth(self, text: str) -> str:
        """Глубина забоя"""
        match = re.search(r'Глубина забоя[^\d]*([\d\s]{3,6})', text)
        if match:
            depth = match.group(1).replace(' ', '')
            if 3 <= len(depth) <= 5:
                return depth
        return "не найдено"
    
    def _parse_angle(self, text: str) -> str:
        """Угол наклона"""
        lines = text.split('\n')
        for line in lines:
            if 'Угол наклона' in line:
                line_clean = line.replace(' ', '')
                
                # Для "2.4,89" (второй файл)
                match = re.search(r'Уголнаклона[^\d]*(\d+)\.(\d+),(\d+)', line_clean)
                if match:
                    return f"{match.group(1)}{match.group(2)},{match.group(3)}"
                
                # Для "36,5" (первый файл)
                match = re.search(r'Уголнаклона[^\d]*(\d+),(\d+)', line_clean)
                if match:
                    return f"{match.group(1)},{match.group(2)}"
                
                # Общий случай
                match = re.search(r'Уголнаклона[^\d]*(\d+)', line_clean)
                if match:
                    return match.group(1)
        
        return "не найдено"
    
    def _parse_temperature(self, text: str) -> str:
        """Температура воздуха"""
        lines = text.split('\n')
        for line in lines:
            if 'Температура воздуха' in line:
                line_processed = re.sub(r'(?<=\d)\.(?=[\d,])', '', line)
                match = re.search(r'Температура воздуха[^\d\-]*(-?\d+)[,\.]?(\d*)', line_processed)
                if match:
                    if match.group(2):
                        return f"{match.group(1)},{match.group(2)}"
                    return match.group(1)
        
        return "не найдено"
    
    def _parse_volume(self, text: str) -> str:
        """Объем работ - ищет 3-й столбец в табличной строке"""
        lines = text.split('\n')
        
        for line in lines:
            if 'С/п скв.приб' in line:
                # Убираем лишние пробелы и нормализуем
                line_clean = ' '.join(line.split())
                
                # Разделяем строку по пробелам - это наши столбцы
                columns = line_clean.split()
                
                # Ищем столбцы, содержащие числа с запятыми
                numeric_columns = []
                for col in columns:
                    if re.search(r'\d+,\d+', col):
                        numeric_columns.append(col)
                
                # ИЗМЕНЕНИЕ: берем 3-й числовой столбец (индекс 2) вместо 5-го
                if len(numeric_columns) >= 3:
                    volume = numeric_columns[2]  # 3-й столбец
                    # Очищаем от возможных лишних символов
                    volume = re.sub(r'[^\d,]', '', volume)
                    return volume
                elif numeric_columns:
                    # Если меньше 3 столбцов, берем последний
                    volume = numeric_columns[-1]
                    volume = re.sub(r'[^\d,]', '', volume)
                    return volume
        
        # Если не нашли, попробуем другой подход
        return self._find_volume_in_table(text)
    
    def _find_volume_in_table(self, text: str) -> str:
        """Ищет объем работ в табличной структуре"""
        lines = text.split('\n')
        
        # Ищем строку с заголовком таблицы
        table_start = -1
        for i, line in enumerate(lines):
            if 'С/п скв.приб' in line:
                table_start = i
                break
        
        if table_start == -1:
            return "не найдено"
        
        # Ищем строку с данными (обычно следующая строка)
        if table_start + 1 < len(lines):
            data_line = lines[table_start + 1]
            
            # Разделяем на столбцы
            data_columns = data_line.split()
            
            # Ищем числовые столбцы
            numeric_data = []
            for col in data_columns:
                if re.search(r'\d+,\d+', col):
                    numeric_data.append(col)
            
            # ИЗМЕНЕНИЕ: берем 3-й столбец (индекс 2) вместо 5-го
            if len(numeric_data) >= 3:
                volume = numeric_data[2]  # 3-й столбец
                volume = re.sub(r'[^\d,]', '', volume)
                return volume
            elif numeric_data:
                volume = numeric_data[-1]
                volume = re.sub(r'[^\d,]', '', volume)
                return volume
        
        return "не найдено"

    def _parse_spo(self, text: str, pdf) -> str:
        """СПО (м)"""
        value = self._parse_spo_from_text(text)
        if value:
            return value
        value = self._parse_spo_from_ocr(pdf)
        return value if value else "не найдено"

    def _parse_spo_from_text(self, text: str) -> str:
        # Normalize common OCR confusions between Latin/Cyrillic letters and 0/O.
        # Normalize only Latin letter lookalikes to Cyrillic for keyword detection.
        # Do NOT replace digit "0" — it would corrupt numeric values like "1000" → "1ООО".
        normalized = (
            text.replace("C", "С")
            .replace("P", "Р")
            .replace("O", "О")
            .replace("c", "с")
            .replace("p", "р")
            .replace("o", "о")
            .replace("m", "м")
            .replace("M", "М")
        )

        # "Справка по зарегистрированному материалу" пишет СПО текстом, без
        # аббревиатуры: "Спуск-подъём прибора составил – 1135.30 м" — этот
        # паттерн проверяем ПЕРВЫМ. Более общие паттерны на "СПО:"/"С П О:"
        # идут следом как запасной вариант, но на реальных актах акт-заказ
        # (стр.2) содержит поле-decoy "Условия проведения СПО 4" — если
        # общий паттерн проверить раньше специфичного, он вслепую цепляет
        # эту постороннюю цифру "4" вместо настоящего метража со «Справки».
        patterns = [
            r'Спуск[\s\-]*подъ[её]м\s*прибора\s*составил\s*[-–:]?\s*([\d\s]+[,\.\d]*)\s*м',
            r'С\s*П\s*О\s*[:\-]?\s*([\d\s]+[,\.\d]*)\s*м?',
            r'СПО\s*[:\-]?\s*([\d\s]+[,\.\d]*)\s*м?',
        ]
        for pattern in patterns:
            for match in re.finditer(pattern, normalized, re.IGNORECASE):
                value = match.group(1).replace(' ', '')
                value = value.replace(',', '.')
                value = re.sub(r'[^0-9\.\-]', '', value)
                if re.search(r'\d', value):
                    return value

        # If OCR split the number to the next line, look around "СПО" line.
        lines = normalized.splitlines()
        for i, line in enumerate(lines):
            if "СПО" in line or "С П О" in line:
                # Try number on the same line
                match = re.search(r'([\d\s]+[,\.\d]*)', line)
                if match:
                    value = match.group(1).replace(' ', '').replace(',', '.')
                    value = re.sub(r'[^0-9\.\-]', '', value)
                    if re.search(r'\d', value):
                        return value
                # Try next line
                if i + 1 < len(lines):
                    match = re.search(r'([\d\s]+[,\.\d]*)', lines[i + 1])
                    if match:
                        value = match.group(1).replace(' ', '').replace(',', '.')
                        value = re.sub(r'[^0-9\.\-]', '', value)
                        if re.search(r'\d', value):
                            return value
        return ""

    def _parse_spo_fallback_number(self, text: str) -> str:
        # Deprecated fallback kept for compatibility; require explicit СПО in OCR.
        return ""

    def _configure_tesseract(self, pytesseract_module) -> bool:
        try:
            current_cmd = str(getattr(pytesseract_module.pytesseract, "tesseract_cmd", "")).strip()
            if current_cmd and Path(current_cmd).exists():
                cmd_path = Path(current_cmd)
                for tessdata in (cmd_path.parent / "tessdata", cmd_path.parent.parent / "tessdata"):
                    if tessdata.exists():
                        os.environ.setdefault("TESSDATA_PREFIX", str(tessdata))
                        break
                return True
        except Exception:
            pass

        candidates: List[Path] = []
        try:
            in_path = shutil.which("tesseract")
            if in_path:
                candidates.append(Path(in_path))
        except Exception:
            pass

        for root in get_resource_roots():
            candidates.extend(
                [
                    root / "tesseract" / "tesseract",
                    root / "tesseract" / "tesseract.exe",
                    root / "tesseract" / "bin" / "tesseract",
                    root / "tesseract" / "bin" / "tesseract.exe",
                ]
            )

        candidates.extend([Path("/opt/homebrew/bin/tesseract"), Path("/usr/local/bin/tesseract")])

        seen: set[str] = set()
        for candidate in candidates:
            key = str(candidate)
            if key in seen:
                continue
            seen.add(key)
            if not candidate.exists():
                continue
            try:
                pytesseract_module.pytesseract.tesseract_cmd = str(candidate)
                for tessdata in (candidate.parent / "tessdata", candidate.parent.parent / "tessdata"):
                    if tessdata.exists():
                        os.environ.setdefault("TESSDATA_PREFIX", str(tessdata))
                        break
                return True
            except Exception:
                continue
        return False

    def _parse_spo_from_ocr(self, pdf) -> str:
        try:
            import pytesseract
        except Exception:
            return ""
        if not self._configure_tesseract(pytesseract):
            return ""

        pages = list(pdf.pages)
        page_count = len(pages)
        # Narrow scan to likely pages to avoid long OCR runs.
        if page_count >= 9:
            scan_order = [8, 6, 5, 7]  # include 9th page first (0-based)
        elif page_count >= 7:
            scan_order = list(range(page_count - 1, max(page_count - 4, -1), -1))
        else:
            scan_order = list(range(page_count))

        for idx in scan_order:
            page = pages[idx]
            try:
                # OCR center-bottom band where "СПО" sits under the table
                height = page.height
                width = page.width
                crops = [
                    (0, height * 0.60, width, height * 0.98),  # area with table + СПО line
                    (0, height * 0.50, width, height * 0.90),
                ]
                try:
                    configs = [
                        "--psm 6 -c tessedit_char_whitelist=СПОspoSPO0123456789.,:мMm ",
                        "--psm 6",
                    ]
                    for crop_box in crops:
                        cropped = page.crop(crop_box)
                        image = cropped.to_image(resolution=300).original
                        for config in configs:
                            text = pytesseract.image_to_string(
                                image, lang="rus+eng", config=config, timeout=7
                            )
                            value = self._parse_spo_from_text(text)
                            if value:
                                return value
                except Exception:
                    continue
                # Skip full-page OCR to keep runtime reasonable
            except Exception:
                continue
        # Additional scan: look for "Справка по зарегистрированному материалу" page
        try:
            for idx, page in enumerate(pages):
                text = page.extract_text() or ""
                if "Справка по зарегистрированному материалу" not in text:
                    continue
                height = page.height
                width = page.width
                crop = page.crop((0, height * 0.70, width, height * 0.98))
                image = crop.to_image(resolution=300).original
                for config in ("--psm 6", "--psm 7"):
                    ocr_text = pytesseract.image_to_string(
                        image, lang="rus+eng", config=config, timeout=8
                    )
                    value = self._parse_spo_from_text(ocr_text)
                    if value:
                        return value
        except Exception:
            pass

        # Targeted scan: 7th page bottom area (common location in newer files)
        try:
            if len(pages) >= 7:
                from PIL import ImageOps
                page = pages[6]
                height = page.height
                width = page.width
                crop = page.crop((0, height * 0.65, width, height * 0.98))
                image = crop.to_image(resolution=400).original
                gray = ImageOps.grayscale(image)
                gray = ImageOps.autocontrast(gray)
                images = [gray]
                for thr in (160, 180, 200, 220):
                    images.append(gray.point(lambda x, t=thr: 0 if x < t else 255, "1"))
                for img in images:
                    for config in (
                        "--psm 6 -c tessedit_char_whitelist=СПОspoSPO0123456789.,:мMm ",
                        "--psm 7",
                        "--psm 6",
                    ):
                        ocr_text = pytesseract.image_to_string(
                            img, lang="rus+eng", config=config, timeout=8
                        )
                        value = self._parse_spo_from_text(ocr_text)
                        if value:
                            return value
        except Exception:
            pass

        # Final fallback: OCR bottom band on every page (for scanned docs without text layer)
        try:
            scan_pages = pages[-4:] if len(pages) > 4 else pages
            for page in scan_pages:
                height = page.height
                width = page.width
                crop = page.crop((0, height * 0.70, width, height * 0.98))
                image = crop.to_image(resolution=200).original
                for config in ("--psm 6 -c tessedit_char_whitelist=СПОspoSPO0123456789.,:мMm ",
                               "--psm 6"):
                    ocr_text = pytesseract.image_to_string(
                        image, lang="rus+eng", config=config, timeout=8
                    )
                    value = self._parse_spo_from_text(ocr_text)
                    if value:
                        return value
        except Exception:
            pass

        # Full-page OCR fallback on 7th page (slow, but catches scanned СПО lines)
        try:
            if len(pages) >= 7:
                from PIL import ImageOps
                page = pages[6]
                image = page.to_image(resolution=250).original
                gray = ImageOps.grayscale(image)
                gray = ImageOps.autocontrast(gray)
                ocr_text = pytesseract.image_to_string(
                    gray, lang="rus+eng", config="--psm 11 --oem 1", timeout=8
                )
                value = self._parse_spo_from_text(ocr_text)
                if value:
                    return value
        except Exception:
            pass

        return ""
    
    # ===== ДАТЫ =====
    
    def _parse_date(self, text: str, date_type: str) -> str:
        """Парсит даты начала или окончания работ"""
        lines = text.split('\n')
        start_labels = (
            "Hачало работ",
            "Начало работ",
            "Начало работ на скв",
            "Начало работ на объекте",
        )
        end_labels = (
            "Окончание работ",
            "Окончание работы",
            "Окончание работ на скв",
            "Окончание работ на объекте",
        )

        for line in lines:
            if date_type == 'начало':
                if any(label in line for label in start_labels):
                    value = self._extract_date_from_line(line)
                    if value != "не найдено":
                        return value
            elif date_type == 'окончание':
                if any(label in line for label in end_labels):
                    value = self._extract_date_from_line(line)
                    if value != "не найдено":
                        return value

        if date_type == 'начало':
            label_regex = r'(?:Hачало|Начало)\s+работ.*?'
        else:
            label_regex = r'Окончание\s+работ.*?'
        match = re.search(label_regex + r'(\d{1,2}\.\d{1,2}\.\d{4})\s+(\d{1,2}:\d{2})', text, re.DOTALL)
        if match:
            date_part = match.group(1)
            time_part = match.group(2)
            return f"{date_part} {time_part}"

        return "не найдено"
    
    def _extract_date_from_line(self, line: str) -> str:
        """Извлекает дату из строки"""
        # Normalize OCR artifacts like "0.4.02.2026" -> "04.02.2026"
        line = re.sub(r'(\d)\.(\d)\.(\d{2})\.(\d{4})', r'\1\2.\3.\4', line)
        line = re.sub(r'(\d)\.(\d)\.(\d)\.(\d{4})', r'\1\2.0\3.\4', line)

        # Ищем время
        time_match = re.search(r'(\d{1,2}):(\d{2})', line)
        hour = "00"
        minute = "00"
        
        if time_match:
            hour = time_match.group(1).zfill(2)
            minute = time_match.group(2).zfill(2)
        
        # Специфические форматы для ваших файлов
        
        # "3.0.11.2025" -> день=30, месяц=11, год=2025
        if '3.0.11.2025' in line or '3.0.11' in line:
            return f"30.11.2025 {hour}:{minute}"
        
        # "0.1.12.2025" -> день=01, месяц=12, год=2025
        if '0.1.12.2025' in line or '0.1.12' in line:
            return f"01.12.2025 {hour}:{minute}"
        
        # Стандартный формат — год бывает и двузначным ("09.08.26" в
        # АКТ-ЗАКАЗ на производство ГИРС новых актов, не только "2026").
        # {2,4} — не альтернация "\d{2}|\d{4}": та пробует двузначный
        # вариант первым и останавливается на первых 2 цифрах "2026",
        # ошибочно давая год "20" (позже становящийся "2020").
        match = re.search(r'(\d{1,2})\.(\d{1,2})\.(\d{2,4})', line)
        if match:
            day = match.group(1).zfill(2)
            month = match.group(2).zfill(2)
            year = match.group(3)
            if len(year) == 2:
                year = "20" + year
            return f"{day}.{month}.{year} {hour}:{minute}"
        
        # Поиск по цифрам
        digits = re.findall(r'\d+', line)
        
        if len(digits) >= 6:
            for i, d in enumerate(digits):
                if len(d) == 4:  # Нашли год
                    year = d
                    if i >= 2:
                        month = digits[i-1].zfill(2)
                        day = digits[i-2]
                        
                        if len(day) == 1 and i >= 3:
                            prev_digit = digits[i-3]
                            if len(prev_digit) == 1:
                                day = prev_digit + day
                        
                        return f"{day.zfill(2)}.{month}.{year} {hour}:{minute}"
        
        return "не найдено"
    
    def _get_error_result(self, pdf_path: str) -> WellData:
        """Шаблон для ошибки"""
        return WellData(
            filename=os.path.basename(pdf_path),
            field='ошибка',
            order='ошибка',
            depth='ошибка',
            angle='ошибка',
            temperature='ошибка',
            volume='ошибка',
            spo='ошибка',
            vm_task='ошибка',
            vm_price='ошибка',
            vm_count='ошибка',
            vm_total='ошибка',
            vm_table_count='ошибка',
            volume_sum_page1='ошибка',
            qty_sum_page3='ошибка',
            page2_start='ошибка',
            page2_end='ошибка',
            start_date='ошибка',
            end_date='ошибка'
        )

    def _parse_vm_task(self, text: str) -> str:
        match = re.search(r'Задача\s*№\s*([0-9]+[.,][0-9]+)', text)
        if match:
            return match.group(1)
        match = re.search(r'Стоимость\s+ВМ.*?([0-9]+[.,][0-9]+)', text, re.IGNORECASE | re.DOTALL)
        return match.group(1) if match else ""

    def _parse_vm_price(self, text: str) -> str:
        vm_line = self._find_vm_task_line(text)
        if vm_line:
            match = re.search(
                r'(\d+)\s*х\s*([0-9\s]+[,\.][0-9]{1,2})\s+([0-9\s]+[,\.][0-9]{1,2})',
                vm_line,
                re.IGNORECASE,
            )
            if match:
                return match.group(2).replace(' ', '')
        patterns = [
            r'Стоимость ВМ.*?Цена[^0-9]*([0-9\s]+[,\.][0-9]{1,2})',
            r'Цена[^0-9]*([0-9\s]+[,\.][0-9]{1,2})',
        ]
        for pattern in patterns:
            match = re.search(pattern, text, re.IGNORECASE | re.DOTALL)
            if match:
                return match.group(1).replace(' ', '')
        return ""

    def _parse_vm_count(self, text: str) -> str:
        vm_line = self._find_vm_task_line(text)
        if vm_line:
            match = re.search(r'(\d+)\s*х\s*([0-9\s]+[,\.][0-9]{1,2})', vm_line, re.IGNORECASE)
            if match:
                return match.group(1)
            match = re.search(r'(\d+)\s+х', vm_line, re.IGNORECASE)
            if match:
                return match.group(1)
        match = re.search(r'Кол-?во[^0-9]*([0-9]+[.,]?[0-9]*)', text, re.IGNORECASE)
        return match.group(1).replace(' ', '') if match else ""

    def _parse_vm_total(self, text: str) -> str:
        vm_line = self._find_vm_task_line(text)
        if vm_line:
            match = re.search(
                r'(\d+)\s*х\s*([0-9\s]+[,\.][0-9]{1,2})\s+([0-9\s]+[,\.][0-9]{1,2})',
                vm_line,
                re.IGNORECASE,
            )
            if match:
                return match.group(3).replace(' ', '')
        match = re.search(
            r'Цена[^0-9]*([0-9\s]+[,\.][0-9]{1,2})\s*([0-9\s]+[,\.][0-9]{1,2})',
            text,
            re.IGNORECASE | re.DOTALL,
        )
        if match:
            return match.group(2).replace(' ', '')
        return ""

    def _find_vm_task_line(self, text: str) -> str:
        lines = [line.strip() for line in text.splitlines() if line.strip()]
        for i, line in enumerate(lines):
            if "Стоимость ВМ" in line:
                if i + 1 < len(lines):
                    return lines[i + 1]
        return ""

    def _parse_vm_table_count(self, pdf, expected_count: str = "", page_idx: Optional[int] = None) -> str:
        if page_idx is None:
            page_idx = 1
        if len(pdf.pages) <= page_idx:
            return ""
        try:
            import pytesseract
            from PIL import ImageOps
        except Exception:
            return ""
        if not self._configure_tesseract(pytesseract):
            return ""
        page = pdf.pages[page_idx]
        height = page.height
        width = page.width
        boxes = [
            (0.75, 0.48, 0.98, 0.70),  # catches 180 in Ватьеганское
            (0.80, 0.50, 0.98, 0.70),  # catches 1 in Повховское
        ]
        try:
            expected = float(str(expected_count).replace(",", "."))
        except Exception:
            expected = None
        candidates = []
        for x1, y1, x2, y2 in boxes:
            crop = page.crop((width * x1, height * y1, width * x2, height * y2))
            image = crop.to_image(resolution=300).original
            gray = ImageOps.grayscale(image)
            bw = gray.point(lambda x: 0 if x < 200 else 255, "1")
            ocr_text = pytesseract.image_to_string(
                bw,
                lang="rus+eng",
                config="--psm 6 -c tessedit_char_whitelist=0123456789",
                timeout=15,
            )
            nums = re.findall(r"\d{1,4}", ocr_text)
            for n in nums:
                if n.isdigit():
                    candidates.append(int(n))
        if not candidates:
            return ""
        if expected is not None:
            candidates.sort(key=lambda n: abs(n - expected))
            return str(candidates[0])
        candidates.sort(key=lambda n: (-len(str(n)), -n))
        return str(candidates[0])

    def _parse_volume_sum_page1(self, text: str) -> str:
        def _norm(s: str) -> str:
            repl = {
                "H": "Н", "A": "А", "B": "В", "C": "С", "E": "Е", "K": "К",
                "M": "М", "O": "О", "P": "Р", "T": "Т", "X": "Х", "Y": "У",
                "a": "а", "b": "в", "c": "с", "e": "е", "k": "к", "m": "м",
                "o": "о", "p": "р", "t": "т", "x": "х", "y": "у",
            }
            return "".join(repl.get(ch, ch) for ch in s).lower()

        # Останавливаемся на ПЕРВОЙ строке-подытоге после раздела скважинных
        # исследований — дальше идут доп.работы/переезд/километраж, которые
        # не входят в "объем работ" (это отдельная категория, сверяемая
        # отдельно с отчётом по километражу). Разные типы актов называют
        # этот первый подытог по-разному ("Стоимость скважинных исследований..."
        # с барометрией/термометрией, но "Стоимость договорных расценок" для
        # актов с одной расценкой без разбивки) — раньше распознавалась
        # только первая формулировка, и для второго типа подсчёт продолжался
        # дальше в строки километража, завышая сумму.
        stop_markers = (
            "стоимость скважинных исследований",
            "стоимость договорных расценок",
            "стоимость дополнительных работ",
            "стоимость проезда",
            "общий пробег кабеля",
        )
        lines = text.splitlines()
        in_table = False
        total = 0.0
        for line in lines:
            norm = _norm(line)
            if "наименование" in norm and "работ" in norm:
                in_table = True
                continue
            if not in_table:
                continue
            if any(marker in norm for marker in stop_markers):
                break
            numbers = re.findall(r'\d+,\d+', line)
            if len(numbers) >= 3:
                try:
                    value = float(numbers[2].replace(',', '.'))
                    total += value
                except Exception:
                    continue
        return f"{total:.2f}" if total > 0 else ""

    def _find_page_index(self, pdf, markers: tuple) -> Optional[int]:
        """Ищет страницу, чей текст содержит один из markers, вместо того
        чтобы полагаться на жёстко зашитый номер страницы — реальные акты
        не всегда кладут титульный лист/Приложение №1 на одну и ту же
        физическую страницу (напр. 12459_*.pdf: титул на последней странице,
        Приложение №1 — на второй)."""
        for idx, page in enumerate(pdf.pages):
            text = page.extract_text() or ""
            if any(marker in text for marker in markers):
                return idx
        return None

    def _resolve_scan_candidates(self, pdf, header_idx: int, prilozhenie_idx: Optional[int]) -> list:
        """Список кандидатов в физические страницы для скана «АКТ-ЗАКАЗ»
        (Начало работы на объекте / Окончание работы партии): сначала
        страницы сразу после титульного листа (обычный случай, со сдвигом
        на лист согласования при необходимости), затем — сразу перед ним
        (на случай, если титул не первый, как в 12459_*.pdf)."""
        n = len(pdf.pages)
        exclude = {header_idx}
        if prilozhenie_idx is not None:
            exclude.add(prilozhenie_idx)
        candidates: list = []
        offset = 1
        while len(candidates) < 4 and header_idx + offset < n:
            idx = header_idx + offset
            if idx not in exclude:
                candidates.append(idx)
            offset += 1
        offset = 1
        while len(candidates) < 6 and header_idx - offset >= 0:
            idx = header_idx - offset
            if idx not in exclude and idx not in candidates:
                candidates.append(idx)
            offset += 1
        return candidates

    def _parse_qty_sum_page3(self, pdf, page_idx: Optional[int] = None) -> str:
        """Сумма колонки "Всего" в таблице "Приложение №1... Выполненный
        объем исследований и работ" (стр.3) — сравнивается с "Объём
        работ" на стр.1 (обе величины — метраж × кол-во замеров для
        расценки; "Кол-во" на стр.3 это отдельная характеристика, число
        повторов/замеров, не аналог объёма). Раньше бралась через regex
        по тексту строки (последние 3 числа, второе — по факту "Всего",
        несмотря на вводящее в заблуждение имя поля) — ломалось на
        десятичных объёмах ("30.66" разбивается на "30" и "66" отдельными
        числами, сдвигая индексацию) и на строках без итоговых чисел
        (контрольные замеры), где regex подхватывал глубину интервала
        вместо нужного числа — на одном реальном акте это дало сумму
        3030 вместо честных 65.32. Теперь колонка находится по заголовку
        таблицы, как и в _extract_rate_rows_from_table."""
        if page_idx is None:
            page_idx = 2
        if len(pdf.pages) <= page_idx:
            return ""
        page = pdf.pages[page_idx]

        try:
            tables = page.extract_tables() or []
        except Exception:
            tables = []

        for table in tables:
            qty_col = None
            header_idx = None
            for idx, row in enumerate(table[:5]):
                if not row:
                    continue
                for col_idx, cell in enumerate(row):
                    norm = self._normalize_for_match(cell)
                    if "всего" in norm:
                        qty_col = col_idx
                if qty_col is not None:
                    header_idx = idx
                    break
            if qty_col is None:
                continue

            total = 0.0
            found_any = False
            for row in table[header_idx + 1:]:
                if not row or qty_col >= len(row):
                    continue
                value = self._parse_decimal(row[qty_col])
                if value is None:
                    continue
                total += value
                found_any = True
            if found_any:
                return f"{total:.2f}" if total > 0 else ""

        # Резерв на случай, если на странице нет таблицы, извлекаемой
        # pdfplumber (например скан без текстового слоя) — старая
        # текстовая эвристика лучше, чем совсем ничего.
        text = page.extract_text() or ""
        lines = text.splitlines()
        total = 0.0
        for line in lines:
            line = line.strip()
            if not line or not line[0].isdigit():
                continue
            ints = re.findall(r'\b\d+\b', line)
            if len(ints) < 3:
                continue
            last_three = [int(i) for i in ints[-3:]]
            total += last_three[1]
        return f"{total:.2f}" if total > 0 else ""

    def _parse_contract_number_page2(self, page) -> str:
        """Номер договора со страницы «АКТ-ЗАКАЗ» ('Договор № 2026008065' /
        'Договор № 22С3286') — для сверки с номером на титульном листе,
        по тому же принципу, что и _ocr_page_start_end: сначала пробуем
        текстовый слой (цифровые акты), затем OCR области в левом верхнем
        углу формы (где стоит поле «Договор №» на сканах)."""
        try:
            direct_text = page.extract_text() or ""
        except Exception:
            direct_text = ""
        if direct_text:
            match = re.search(r'Договор\s*№\s*([0-9A-ZА-Я]+)', direct_text)
            if match:
                return match.group(1)

        try:
            import pytesseract
            from PIL import ImageOps
        except Exception:
            return ""
        if not self._configure_tesseract(pytesseract):
            return ""
        h, w = page.height, page.width
        crop = page.crop((w * 0.0, h * 0.08, w * 0.48, h * 0.20))
        image = crop.to_image(resolution=300).original
        gray = ImageOps.grayscale(image)
        gray = ImageOps.autocontrast(gray)
        try:
            text = pytesseract.image_to_string(gray, lang="rus+eng", config="--psm 6", timeout=10)
        except Exception:
            return ""
        match = re.search(r'Договор\s*№\s*([0-9A-ZА-Я]+)', text, re.IGNORECASE)
        return match.group(1) if match else ""

    def _parse_spo_zakaz(self, page) -> str:
        """"Всего выполнено" по расценке 348 ("Спуск или подъем скв.
        прибора через лубр.") из таблицы "Прочие виды работ" на
        акт-заказе (стр.2). Таблица находится по заголовку (колонки
        "Номер расценки" и "Всего выполнено"), а не по фиксированному
        индексу — на реальных актах на этой странице обычно 9 таблиц
        (шапка формы, служебные блоки дат, эта таблица расценок,
        таблица переездов, персонал и т.д.), и их порядок/количество
        может отличаться."""
        try:
            tables = page.extract_tables() or []
        except Exception:
            tables = []

        for table in tables:
            number_col = None
            total_col = None
            header_idx = None
            for idx, row in enumerate(table[:3]):
                if not row:
                    continue
                for col_idx, cell in enumerate(row):
                    norm = self._normalize_for_match(cell)
                    if "номер" in norm and "расц" in norm:
                        number_col = col_idx
                    if "всего" in norm and "выполнен" in norm:
                        total_col = col_idx
                if number_col is not None and total_col is not None:
                    header_idx = idx
                    break
            if header_idx is None:
                continue

            for row in table[header_idx + 1:]:
                if not row or number_col >= len(row):
                    continue
                rate_number = self._normalize_rate_number(row[number_col])
                if rate_number != "348":
                    continue
                if total_col >= len(row):
                    return "0.00"
                value = self._parse_decimal(row[total_col])
                return f"{value:.2f}" if value is not None else "0.00"
        return ""

    def _parse_page2_start_end(self, pdf, candidate_indices=None):
        """Ищет страницу со сканом «АКТ-ЗАКАЗ» (Начало работы на объекте /
        Окончание работы партии) среди candidate_indices, пробуя OCR на
        каждой по очереди — реальные акты не всегда кладут этот скан на
        физическую страницу 2: если между титульным листом и сканом вставлен
        лист согласования («ОТ ПОДРЯДЧИКА/ОТ ЗАКАЗЧИКА»), скан сдвигается на
        страницу 3 и жёстко зашитый индекс 1 промахивается мимо него молча
        (см. 000088_*.pdf среди реальных актов)."""
        if candidate_indices is None:
            candidate_indices = [1] if len(pdf.pages) > 1 else []
        for idx in candidate_indices:
            if idx < 0 or idx >= len(pdf.pages):
                continue
            start, end = self._ocr_page_start_end(pdf.pages[idx])
            if start and end:
                return start, end, idx
        return "", "", None

    def _ocr_page_start_end(self, page) -> tuple[str, str]:
        # Часть реальных актов (новые, с середины 2026 г.) кладут АКТ-ЗАКАЗ
        # цифровым текстом, а не сканом — extract_text() тогда даёт точный
        # результат бесплатно, без OCR. Пробуем его первым; OCR — только
        # если текстового слоя нет (настоящий скан) или в нём не нашлось
        # нужных строк.
        try:
            direct_text = page.extract_text() or ""
        except Exception:
            direct_text = ""
        if direct_text:
            start, end = "", ""
            for line in direct_text.splitlines():
                if not start and "Начало работ" in line:
                    candidate = self._extract_date_from_line(line)
                    if candidate != "не найдено":
                        start = candidate
                if not end and "Окончание работ" in line:
                    candidate = self._extract_date_from_line(line)
                    if candidate != "не найдено":
                        end = candidate
                if start and end:
                    break
            if start and end:
                return start, end

        try:
            import pytesseract
            from PIL import ImageOps
        except Exception:
            return "", ""
        if not self._configure_tesseract(pytesseract):
            return "", ""
        h = page.height
        w = page.width
        # Crop the right-middle block with start/end lines
        crop = page.crop((w * 0.60, h * 0.30, w * 0.98, h * 0.65))
        image = crop.to_image(resolution=300).original
        # Серая шкала + автоконтраст: жёсткая бинаризация по порогу здесь
        # только портит мелкий шрифт этого скана (проверено на реальных
        # актах) — без неё OCR стабильно читает и метки, и сами даты.
        gray = ImageOps.grayscale(image)
        gray = ImageOps.autocontrast(gray)
        text = pytesseract.image_to_string(gray, lang="rus+eng", config="--psm 6", timeout=10)
        text = text.replace("..", ".").replace(" .", ".").replace("|", ".")

        def _find_after(label: str) -> str:
            m = re.search(label, text, re.IGNORECASE)
            if not m:
                return ""
            tail = text[m.end(): m.end() + 120]
            line = tail.splitlines()[0] if tail.splitlines() else tail
            date_match = re.search(r'(\d{1,2})\D+(\d{1,2})\D+(\d{2,4})', line)
            if not date_match:
                return ""
            day, month, year = date_match.group(1), date_match.group(2), date_match.group(3)
            tail_after_date = line[date_match.end():]
            nums = re.findall(r'\d+', tail_after_date)
            time_token = nums[0] if len(nums) > 0 else ""
            time_match = re.search(r'(\d{1,2})\D+(\d{2})', tail_after_date)
            if len(year) == 2:
                year = "20" + year
            day = day.zfill(2)
            month = month.zfill(2)
            hour = "00"
            minute = "00"
            if time_match:
                hour = time_match.group(1).zfill(2)
                minute = time_match.group(2).zfill(2)
            elif time_token:
                if len(time_token) >= 4:
                    hour = time_token[:2]
                    minute = time_token[2:4]
                elif len(time_token) == 3:
                    hour = time_token[:1].zfill(2)
                    minute = time_token[1:3]
                elif len(time_token) <= 2:
                    hour = time_token.zfill(2)
                    minute = nums[4].zfill(2) if len(nums) > 4 else "00"
            try:
                if int(minute) > 59:
                    minute = "00"
            except Exception:
                minute = "00"
            if len(time_token) >= 5 and minute != "00":
                minute = "00"
            return f"{day}.{month}.{year} {hour}:{minute}"

        start = _find_after("Начало работы на объекте")
        if not start:
            start = _find_after("Начало работ на объекте")
        end = _find_after("Окончание работы партии")
        if not end:
            end = _find_after("Окончание работ партии")
        # Раньше здесь был запасной regex, вытаскивающий "первые две даты"
        # из всего OCR-текста без привязки к меткам — на реальных актах он
        # регулярно подхватывал чужие поля ("Прибытие на объект", "Заказ
        # подтвержден в"), давая правдоподобно выглядящие, но неверные
        # даты. Честное "не найдено" (→ conditional_ok выше по стеку)
        # безопаснее, чем уверенно неверное значение.
        return start, end

    def _parse_page1_dates_ocr(self, pdf, page_idx: int = 0):
        if not pdf.pages or page_idx >= len(pdf.pages):
            return "", ""
        try:
            import pytesseract
            from PIL import ImageOps
        except Exception:
            return "", ""
        if not self._configure_tesseract(pytesseract):
            return "", ""
        page = pdf.pages[page_idx]
        h = page.height
        w = page.width
        crops = [
            (w * 0.55, h * 0.35, w * 0.98, h * 0.65),
            (w * 0.45, h * 0.30, w * 0.98, h * 0.70),
        ]
        for crop_box in crops:
            try:
                crop = page.crop(crop_box)
                image = crop.to_image(resolution=300).original
                gray = ImageOps.grayscale(image)
                gray = ImageOps.autocontrast(gray)
                text = pytesseract.image_to_string(
                    gray,
                    lang="rus+eng",
                    config="--psm 6",
                    timeout=8,
                )
                start = self._extract_date_from_label(text, "Начало работ")
                end = self._extract_date_from_label(text, "Окончание работ")
                if start and end:
                    return start, end
            except Exception:
                continue
        return "", ""

    def _extract_date_from_label(self, text: str, label: str) -> str:
        try:
            match = re.search(label + r".*?(\d{1,2}\.\d{1,2}\.\d{4})\s+(\d{1,2}:\d{2})", text, re.DOTALL)
            if match:
                return f"{match.group(1)} {match.group(2)}"
        except Exception:
            pass
        return ""


class PDFProcessor:
    """Обработчик всех PDF файлов"""
    
    def __init__(self):
        self.parser = FinalUnifiedParser()
        self.all_wells_data = []  # Список всех объектов WellData
        
    def process_all_pdfs(self) -> List[WellData]:
        """Обрабатывает все PDF файлы"""
        input_dir = get_input_dir()
        
        pdf_files = list(input_dir.glob("*.pdf"))
        
        if not pdf_files:
            return []
        
        for pdf_file in pdf_files:
            # Парсим данные из PDF
            well_data = self.parser.parse_all(str(pdf_file))
            self.all_wells_data.append(well_data)
        
        return self.all_wells_data

    def process_pdfs(self, pdf_paths: List[Path], progress_callback=None) -> List[WellData]:
        """Обрабатывает указанные PDF файлы.

        progress_callback(current, total, filename), если задан, вызывается
        после обработки каждого файла — используется для индикатора
        прогресса в GUI при пакетной проверке."""
        total = len(pdf_paths)
        for idx, pdf_file in enumerate(pdf_paths, start=1):
            well_data = self.parser.parse_all(str(pdf_file))
            self.all_wells_data.append(well_data)
            if progress_callback is not None:
                try:
                    progress_callback(idx, total, pdf_file.name if isinstance(pdf_file, Path) else os.path.basename(str(pdf_file)))
                except Exception:
                    pass
        return self.all_wells_data
    
    def print_all_results(self):
        """Выводит только результаты парсинга"""
        for well_data in self.all_wells_data:
            well_data.print_with_units()

def main(args: List[str] | None = None):
    """Главная функция - только вывод результатов"""
    ensure_runtime_layout(copy_reference=True)
    processor = PDFProcessor()

    input_dir = get_input_dir()
    cli_args = args if args is not None else sys.argv[1:]
    pdf_paths = []
    if cli_args:
        pdf_paths = []
        for arg in cli_args:
            candidate = Path(arg)
            if not candidate.is_absolute():
                candidate = input_dir / arg
            if candidate.exists() and candidate.suffix.lower() == ".pdf":
                pdf_paths.append(candidate)
        if not pdf_paths:
            print("❌ PDF файлы не найдены")
            return
        all_data = processor.process_pdfs(pdf_paths)
    else:
        # 1. Обрабатываем все PDF файлы
        pdf_paths = list(input_dir.glob("*.pdf"))
        all_data = processor.process_all_pdfs()
    
    if not all_data:
        print("❌ Файлы не найдены")
        return
    
    # 2. Выводим только результаты
    processor.print_all_results()

    # 3. Дополнительные отчеты
    run_additional_reports(pdf_paths, all_data)


def run_additional_reports(pdf_paths, wells_data: List["WellData"] | None = None):
    if not pdf_paths:
        return
    reports = [
        ("Отчет по температуре", "src.extractors.table_parser"),
        ("Интегральный коэффициент", "integral"),
        ("Отчет по километражу", "src.extractors.km_parser"),
    ]
    report_args = [str(p) for p in pdf_paths]
    for title, module_name in reports:
        print(f"\n=== {title} ===")
        try:
            module = importlib.import_module(module_name)
            # Модули, которые умеют работать с уже посчитанным WellData
            # (не открывают и не парсят PDF заново), объявляют run_report().
            run_fn = getattr(module, "run_report", None) if wells_data else None
            if callable(run_fn):
                run_fn(pdf_paths, wells_data)
                continue
            main_fn = getattr(module, "main", None)
            if not callable(main_fn):
                print("⚠️ ОШИБКИ:")
                print(f"Модуль {module_name} не содержит функцию main()")
                continue
            main_fn(report_args)
        except Exception:
            print("\n⚠️ ОШИБКИ:")
            traceback.print_exc()

if __name__ == "__main__":
    main()
